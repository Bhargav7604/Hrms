"""IP chart — the per-admission inpatient workspace and the screens around it.

* New Admission: patient lookup (UHID / name / Aadhaar last 4 / mobile), payer category and sponsor,
  credit-letter and approved amount, primary + second consultant, ward-filtered vacant bed, package.
  Admissions are created through the record service, so adt.py's bed and duplicate-stay rules apply.
* Admissions list with KPI filters, ward acceptance ("Accept patient"), clinical condition with a
  who/when log, patient allergy status (NKA / nine types / substance → pharmacy safety checks) and
  reactive (infection-precaution) status that raises the admission's isolation level.
* Chart data: header banner, vitals tiles with range flags + NEWS2 + trend, "Chart it" observations,
  recent investigations / medicines / procedures / notes, patient visit tracker, timeline, progress
  notes by shift, doctor's instructions with acknowledgement, structured consents.
* Order Desk: investigations with line discounts grouped into indents (collect payment / covered by
  advance), service requests, medication indents with stock availability and safety screening,
  pharmacy returns, procedure / grade / manual-equipment charges, equipment connect → disconnect
  auto-billing by the machine's rate rule, voids with reason and voided-by audit, plus the
  medicines-on-this-stay and investigations ledgers.
* On Bed: every occupied bed with stay, consultant, acuity and bill / collected / balance.
* Discharge tracker: Admitted → Discharge initiated → Clearances → Settlement → Discharged, with
  departmental clearances signed off by each department's role and configured roles notified per step.

Discharge settings live in the settings store (namespace "discharge"):
    {"clearances": ["Pharmacy returns", "Nursing", "Billing", ...],
     "notify": {"admitted": ["nurse"], "initiated": ["billing", "pharmacist"], ...}}
"""

from __future__ import annotations

import csv
import io
import logging
import math
import re
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Any, Literal

import sqlalchemy as sa
from fastapi import APIRouter, Request
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

from ..config import settings
from ..core import records as rs
from ..core.audit import audit
from ..core.hooks import emit, hook
from ..core.security import CtxDep
from ..core.sequences import next_code
from ..core.util import (
    ApiError,
    age_from,
    bad_request,
    conflict,
    forbidden,
    is_uuid,
    minutes_between,
    money,
    new_id,
    not_found,
    now,
    parse_dt,
    to_json_value,
)
from ..modules.clinical.ipd import CONDITIONS, ISOLATION
from ..schema import (
    ENTITY_TABLES,
    beds,
    charges,
    employees,
    encounters,
    invoices,
    orders,
    patients,
    payments,
    records,
    services,
    stock_moves,
)
from ._common import E, create, find, first, get, has, live, transition, update
from .billing import clearance, deposit_balance, post_charge
from .inventory import item_label, items_by_id
from .pharmacy import active_medications, patient_allergies, safety_check

try:  # layered settings store (core); absent → built-in defaults
    from ..core.config_store import get_settings as _get_settings
except ImportError:  # pragma: no cover
    _get_settings = None
try:  # field configuration for custom screens
    from .fieldconfig import check_screen as _check_screen
except ImportError:  # pragma: no cover
    _check_screen = None

log = logging.getLogger("vlack.ipchart")
router = APIRouter(prefix="/api/ipchart", tags=["ipchart"])

ADM, BED, WARDS, DEPTS, PATIENTS = "ipd.admissions", "ipd.beds", "mdm.wards", "mdm.departments", "registration.patients"
EQUIP, USAGE, INSTR, TRF = "ipd.equipment", "ipd.equipment_usage", "ipd.instructions", "ipd.transfers"
CLR, CHK, SUMMARY = "discharge.clearances", "discharge.checklists", "discharge.summaries"
VITALS, NOTES = "emr.vitals", "emr.notes"
ORD, INDENT, DISPENSE, RETURN = "orders.orders", "pharmacy.indents", "pharmacy.dispenses", "pharmacy.returns"
CONSENT, CONSENT_TPL, DOCS = "consents.consents", "consents.templates", "documents.documents"
PAYERS, PACKAGES, CHARGES, PAYMENTS = "insurance.payers", "billing.packages", "billing.charges", "billing.payments"
ASSESS, TASKS, MAR = "nursing.assessments", "nursing.tasks", "nursing.mar"
CONSULTS, OT_BOOKINGS = "caremgmt.consults", "ot.bookings"

OPEN = ("Admitted", "Discharge Initiated")
TS = sa.DateTime(timezone=True)
PAYER_LABEL = {"Cash": "General", "Insurance": "Insurance", "Corporate": "Corporate", "Government Scheme": "Govt. scheme"}
SPONSOR_TYPES = {"Insurance": {"Insurer", "TPA"}, "Corporate": {"Corporate"}, "Government Scheme": {"Government Scheme"}}
ADMISSION_TYPES = ["Elective", "Emergency", "Day Care", "Transfer-in", "Maternity"]
ADMISSION_SOURCES = ["OPD", "ER", "Direct", "Referral", "OT"]

ALLERGY_TYPES = ["Drug", "Food", "Contact", "Dust", "Pollen", "Mould", "Latex", "Insect bite", "Skin"]
REACTIVE = ["HIV", "Hepatitis B", "Hepatitis C", "Tuberculosis", "COVID-19", "Influenza", "Chicken pox", "Whooping cough",
            "MDRO", "Neutropenia"]
# Transmission-based precautions a reactive status implies (standard precautions for blood-borne viruses).
ISOLATION_FOR = {"Tuberculosis": "Airborne", "COVID-19": "Airborne", "Chicken pox": "Airborne", "Influenza": "Droplet",
                 "Whooping cough": "Droplet", "MDRO": "Contact", "Neutropenia": "Protective (neutropenic)"}
ISOLATION_RANK = {"None": 0, "Contact": 1, "C. difficile (spore)": 2, "Protective (neutropenic)": 2, "Droplet": 3, "Airborne": 4}

SHIFTS = ["Morning", "Evening", "Night"]
STEPS = [("admitted", "Admitted"), ("initiated", "Discharge initiated"), ("clearances", "Clearances"),
         ("settlement", "Settlement"), ("discharged", "Discharged")]
STEP_KEYS = [k for k, _ in STEPS]
DEFAULT_CLEARANCES = ["Pharmacy returns", "Nursing", "Billing", "Laboratory (pending reports)", "Radiology (films / CDs)", "Dietary"]
# Who signs off each department's clearance (admin may sign any).
CLEARANCE_ROLES: dict[str, set[str]] = {
    "Pharmacy returns": {"pharmacist"}, "Billing": {"billing"}, "Insurance / TPA": {"billing"}, "Nursing": {"nurse"},
    "Laboratory (pending reports)": {"lab"}, "Radiology (films / CDs)": {"radiology"}, "Dietary": {"nurse", "support"},
    "Housekeeping / linen": {"support"}, "Medical records": {"front_office", "quality"},
}
# A cleared department ticks the matching discharge-checklist item (adt.py signs off by role too).
CLEARANCE_CHECKLIST = {"Pharmacy returns": "pharmacy_returns_done", "Nursing": "nursing_clearance", "Billing": "billing_cleared"}
DEFAULT_NOTIFY: dict[str, list[str]] = {"admitted": ["nurse"], "initiated": ["nurse", "billing", "pharmacist"],
                                        "clearances": ["billing"], "settlement": ["nurse", "front_office"], "discharged": ["support"]}
STAFF_ROLES = ["doctor", "nurse", "front_office", "billing", "pharmacist", "lab", "radiology", "support", "quality", "management"]

INVESTIGATION_CATS = ("LAB", "RAD", "PROC")
SERVICE_REQUEST_CATS = {"NURSING": "NURSING", "DIET": "DIET", "PHYSIO": "PROC", "DIALYSIS": "PROC", "DENTAL": "PROC"}
CHARGE_CATS = ("PROC", "OT", "ICU", "CONSULT", "NURSING", "MISC", "BLOOD", "AMBULANCE", "PHYSIO", "DIALYSIS", "DENTAL")
GRADE_CATS = ("CONSULT", "NURSING", "ICU", "ROOM", "PROC", "PHYSIO")
INVESTIGATION_TO_ORDER = {"LAB": "LAB", "RAD": "RAD", "PROC": "PROC"}
PHARM_SOURCES = (DISPENSE, INDENT, RETURN)
INDENT_STORES = ["IP Pharmacy"]  # ward indents are issued by the IP pharmacy (pharmacy.py)
INDENT_CATEGORIES = {"Normal": "Routine", "Urgent": "Urgent", "STAT": "STAT", "Discharge medication": "Routine"}
GRADES = {"General": "rate", "Private": "rate_private", "ICU": "rate_icu", "International": "rate_international"}
PRIVATE_BEDS = {"Semi-Private", "Private", "Deluxe", "Suite"}
ICU_BEDS = {"ICU", "HDU", "NICU", "PICU"}

# Vital-sign bands (adult): (critical_low, low, high, critical_high) — value ≤ critical_low → Critical low,
# < low → Low, ≥ critical_high → Critical high, > high → High, else Normal.
RANGES: dict[str, tuple[float | None, float | None, float | None, float | None]] = {
    "bp_systolic": (89, 100, 139, 180), "bp_diastolic": (39, 60, 89, 120), "pulse": (40, 51, 100, 131),
    "resp_rate": (8, 12, 20, 25), "spo2": (91, 95, None, None), "temperature_c": (35.0, 36.1, 38.0, 39.1),
    "blood_sugar": (54, 70, 180, 300), "pain_score": (None, None, 3, 7), "bmi": (None, 18.5, 24.9, 40),
}
VITAL_LABELS = {"bp_systolic": "Sys", "bp_diastolic": "Dia", "pulse": "Pulse", "resp_rate": "RR", "spo2": "SpO₂",
                "temperature_c": "Temp", "blood_sugar": "GRBS", "pain_score": "Pain", "bmi": "BMI"}
SEVERITY = {"Normal": 0, "Low": 1, "High": 1, "Critical low": 2, "Critical high": 2}

NOTE_TEMPLATES: dict[str, list[dict[str, str]]] = {
    "clinical": [
        {"name": "SOAP progress note", "text": "S: \nO: Vitals — BP __/__, HR __, RR __, SpO₂ __%, Temp __ °C. Examination: \nA: \nP: "},
        {"name": "Daily ward round", "text": "Day __ of admission. Complaints: \nOvernight events: \nExamination: \nInvestigations reviewed: \n"
                                             "Assessment: \nPlan: continue / modify treatment; diet; activity; review."},
        {"name": "Post-operative day", "text": "POD __ after ____. Pain score __/10. Drain __ ml. Wound: clean / soaked. Bowel: \n"
                                               "Plan: mobilise, oral fluids, analgesia, drain removal when < 50 ml/24 h."},
        {"name": "Critical care review", "text": "Haemodynamics: \nVentilation / O₂: \nNeuro (GCS __): \nRenal (UO __ ml/kg/h): \n"
                                                 "Infection: \nLines / tubes: \nPlan: "},
        {"name": "Discharge planning", "text": "Clinically fit for discharge on __. Pending: \nMedicines reconciled: \nFollow-up: "},
    ],
    "nursing": [
        {"name": "Shift summary", "text": "Patient conscious / oriented. Vitals stable. Intake __ ml, output __ ml. Medicines given as charted. "
                                          "Pain __/10. Safety: bed rails up, call bell in reach. Concerns: "},
        {"name": "Wound / dressing", "text": "Dressing done at __. Wound clean / discharge ____. Patient tolerated well."},
        {"name": "Fall-risk precautions", "text": "Morse __. Fall-risk band applied, bed lowest, rails up, attendant educated."},
        {"name": "Blood transfusion", "text": "Unit no. __, started at __, rate __ ml/h. Pre-vitals stable. No reaction observed at 15 min."},
    ],
}

GENERAL_CONSENT_TEXT = (
    "I, {patient_name} (UHID {uhid}), or my authorised representative, consent to admission to {hospital} and to the "
    "examination, investigations, nursing care and routine treatment advised by the treating doctor {doctor}. The nature of "
    "my illness, the proposed plan of care, its expected benefits, the common risks and the alternatives, including the option "
    "of refusing treatment, have been explained to me in a language I understand. I may ask questions and withdraw this consent "
    "at any time. Separate written consent will be taken for any surgery, anaesthesia, blood transfusion or high-risk procedure. "
    "I agree to follow hospital rules and to pay for services as per the hospital tariff. Date: {date}."
)


# ================================================================== small helpers

def _stamp() -> str:
    return now().isoformat(timespec="seconds")


def _f(v: Any) -> float | None:
    try:
        return None if v in (None, "") else float(v)
    except (TypeError, ValueError):
        return None


def _name(first_name: str | None, last_name: str | None) -> str:
    return " ".join(x for x in (first_name, last_name) if x) or "—"


def _emp(ctx) -> str | None:
    return ctx.user.get("employee_id")


def _who(ctx) -> str:
    return ctx.user.get("name") or ctx.role


def _can(ctx, *roles: str) -> bool:
    return ctx.role == "admin" or ctx.role in roles


# One table of who may do what on the IP screens. Endpoints enforce it (`_need`) and the chart
# header sends the same answers to the client, which renders only the controls a role can use.
PERMS: dict[str, tuple[tuple[str, ...], str | None, str]] = {
    # key: (roles allowed by the workflow, entity the role matrix must grant, level needed)
    "admit": (("front_office", "nurse", "doctor"), ADM, "write"),
    "accept": (("nurse",), ADM, "write"), "condition": (("doctor", "nurse"), ADM, "write"),
    "clinical": (("doctor", "nurse"), ADM, "write"), "observe": (("doctor", "nurse"), VITALS, "write"),
    "order": (("doctor", "nurse", "front_office"), ORD, "write"), "void_order": (("doctor", "nurse", "front_office"), ORD, "write"),
    "indent": (("doctor", "nurse"), INDENT, "write"), "void_indent": (("doctor", "nurse", "pharmacist"), INDENT, "write"),
    "pharmacy_return": (("nurse", "pharmacist"), RETURN, "write"),
    "charge": (("doctor", "nurse", "front_office", "billing"), ADM, "read"),
    "equipment": (("doctor", "nurse"), ADM, "write"), "void_usage": (("doctor", "nurse", "billing"), ADM, "read"),
    "collect": (("billing", "front_office"), PAYMENTS, "write"), "discount": (("billing", "front_office"), CHARGES, "write"),
    "rate_override": (("billing",), CHARGES, "write"), "void_charge": (("billing", "front_office"), CHARGES, "write"),
    "money": (("billing", "front_office", "management"), CHARGES, "read"),
    "discharge": (("doctor", "nurse"), ADM, "write"), "revoke_discharge": (("doctor",), ADM, "write"),
    "sponsor": (("front_office", "billing"), ADM, "read"), "change_patient": (("front_office",), ADM, "write"),
    "doctor_note": (("doctor",), NOTES, "write"), "nursing_note": (("nurse",), NOTES, "write"), "instruct": (("doctor",), INSTR, "write"),
    "consent": (("doctor", "nurse", "front_office"), CONSENT, "write"), "documents": (("doctor", "nurse", "front_office"), DOCS, "write"),
    "configure": (("doctor", "nurse", "billing", "front_office"), ADM, "read"),
}
CLINICAL_VIEW = ("doctor", "nurse", "quality", "management")


def _perm(ctx, key: str) -> bool:
    cache = ctx.__dict__.setdefault("_ipchart_perms", {})
    if key not in cache:
        roles, fqn, mode = PERMS[key]
        cache[key] = _can(ctx, *roles) and (fqn is None or _ent_ok(ctx, fqn, mode))
    return cache[key]


def _need(ctx, key: str, what: str = "do this") -> None:
    if not _perm(ctx, key):
        raise forbidden(f"Your role cannot {what}")


def _ent_ok(ctx, fqn: str, mode: str = "read") -> bool:
    """Entity access as the role matrix grants it (no side effects)."""
    if not has(fqn):
        return False
    try:
        rs.ensure_access(ctx, E(fqn), mode)
        return True
    except ApiError:
        return False


def _money_ok(ctx) -> bool:
    """Amounts (bill, collected, balance, rates) are shown to billing / front office / management only."""
    return _perm(ctx, "money")


def _dcol(fqn: str, key: str):
    return sa.cast(E(fqn).table.c.data[key].astext, TS)


def _iso(v: Any) -> str | None:
    return to_json_value(v) if v is not None else None


async def _rows(ctx, fqn: str, *conds, order=None, limit: int = 500) -> list[dict]:
    if not has(fqn):
        return []
    ent = E(fqn)
    stmt = sa.select(ent.table).where(*rs.base_where(ent), *conds)
    if order is not None:
        stmt = stmt.order_by(*(order if isinstance(order, (list, tuple)) else [order]))
    return [rs.to_record(r) for r in (await ctx.conn.execute(stmt.limit(limit))).all()]


async def _emp_names(ctx, ids) -> dict[str, str]:
    want = [i for i in {str(x) for x in ids if x} if is_uuid(i)]
    if not want:
        return {}
    rows = (await ctx.conn.execute(sa.select(employees.c.id, employees.c.name).where(employees.c.id.in_(want)))).all()
    return {str(i): n for i, n in rows}


async def _user_names(ctx, ids) -> dict[str, str]:
    from ..schema import users

    want = [i for i in {str(x) for x in ids if x} if is_uuid(i)]
    if not want:
        return {}
    rows = (await ctx.conn.execute(sa.select(users.c.id, users.c.name).where(users.c.id.in_(want)))).all()
    return {str(i): n for i, n in rows}


async def _admission(ctx, aid: str) -> dict:
    adm = await get(ctx, ADM, aid) if is_uuid(aid) else None
    if not adm:
        raise not_found("Admission not found")
    return adm


def _require_open(adm: dict, what: str = "change this admission") -> None:
    if adm.get("status") not in OPEN:
        raise conflict(f"Admission {adm.get('code')} is {str(adm.get('status')).lower()} — cannot {what}")


async def _patient_row(ctx, pid: str) -> dict:
    row = (await ctx.conn.execute(sa.select(patients).where(patients.c.id == pid))).first() if is_uuid(pid) else None
    if not row:
        raise not_found("Patient not found")
    p = rs.to_record(row)
    p["age"] = age_from(p.get("dob"))
    return p


def _patient_brief(p: dict) -> dict:
    idn = str(p.get("id_number") or "")
    return {"id": p["id"], "uhid": p.get("code"), "name": p.get("title") or _name(p.get("first_name"), p.get("last_name")),
            "gender": p.get("gender"), "age": p.get("age", age_from(p.get("dob"))), "dob": p.get("dob"), "phone": p.get("phone"),
            "blood_group": p.get("blood_group"), "aadhaar_last4": idn[-4:] if len(idn) >= 4 else None, "status": p.get("status"),
            "payer_type": p.get("payer_type")}


def _los_days(started: Any, ended: Any = None) -> int | None:
    s = parse_dt(started)
    if not s:
        return None
    e = parse_dt(ended) or now()
    return max(0, (e.astimezone(settings.tz).date() - s.astimezone(settings.tz).date()).days)


def vital_status(key: str, value: Any) -> str | None:
    v = _f(value)
    band = RANGES.get(key)
    if v is None or not band:
        return None
    cl, lo, hi, ch = band
    if cl is not None and v <= cl:
        return "Critical low"
    if ch is not None and v >= ch:
        return "Critical high"
    if lo is not None and v < lo:
        return "Low"
    if hi is not None and v > hi:
        return "High"
    return "Normal"


def _worst(*statuses: str | None) -> str | None:
    vals = [s for s in statuses if s]
    return max(vals, key=lambda s: SEVERITY.get(s, 0)) if vals else None


def _temp_c(v: dict) -> float | None:
    if _f(v.get("temperature_c")) is not None:
        return _f(v.get("temperature_c"))
    t = _f(v.get("temperature"))
    return round((t - 32) * 5 / 9, 1) if t is not None else None


def _vital_point(v: dict) -> dict:
    tc = _temp_c(v)
    p = {"id": v["id"], "recorded_at": v.get("recorded_at"), "bp_systolic": _f(v.get("bp_systolic")), "bp_diastolic": _f(v.get("bp_diastolic")),
         "pulse": _f(v.get("pulse")), "resp_rate": _f(v.get("resp_rate")), "spo2": _f(v.get("spo2")), "temperature_c": tc,
         "blood_sugar": _f(v.get("blood_sugar")), "pain_score": _f(v.get("pain_score")), "bmi": _f(v.get("bmi")),
         "weight": _f(v.get("weight")), "o2_support": v.get("o2_support"), "gcs": _f(v.get("gcs")), "news2": v.get("news2"),
         "news2_risk": v.get("news2_risk"), "recorded_by": v.get("recorded_by")}
    p["status"] = {k: vital_status(k, p.get(k)) for k in RANGES}
    p["status"]["bp"] = _worst(p["status"]["bp_systolic"], p["status"]["bp_diastolic"])
    return p


def _out_of_range(point: dict) -> list[dict]:
    out = []
    for k, label in VITAL_LABELS.items():
        s = point["status"].get(k)
        if s and s != "Normal":
            out.append({"key": k, "label": label, "value": point.get(k), "status": s})
    return out


# ================================================================== settings (discharge)

async def discharge_settings(ctx) -> dict:
    defaults = {"clearances": DEFAULT_CLEARANCES, "notify": DEFAULT_NOTIFY}
    vals: dict = {}
    if _get_settings is not None:
        try:
            vals = await _get_settings(ctx, "discharge", {}) or {}
        except Exception:  # settings table missing / not migrated yet → defaults
            log.debug("discharge settings unavailable", exc_info=True)
            vals = {}
    allowed = set(E(CLR).field_map["department_type"].options or [])
    clr = [c for c in (vals.get("clearances") or defaults["clearances"]) if c in allowed] or DEFAULT_CLEARANCES
    raw = vals.get("notify") or vals.get("notify_roles") or {}
    notify = {k: list(DEFAULT_NOTIFY.get(k, [])) for k in STEP_KEYS}
    if isinstance(raw, dict):
        for k, roles in raw.items():
            if k in notify and isinstance(roles, list):
                notify[k] = [r for r in roles if r in STAFF_ROLES]
    return {"clearances": clr, "notify": notify}


# ================================================================== patient lookup & admission form

@router.get("/patients/search")
async def patient_search(ctx: CtxDep, uhid: str | None = None, name: str | None = None, aadhaar: str | None = None,
                         mobile: str | None = None, q: str | None = None, limit: int = 20):
    """New Admission lookup. Minimums: UHID 4 characters, name 3 letters, Aadhaar the last 4 digits,
    mobile 5 digits. Fields combine with AND; `q` (a scan) matches a UHID or IP number exactly."""
    rs.ensure_access(ctx, E(ADM), "read")
    t = patients
    conds = [t.c.deleted_at.is_(None), sa.func.coalesce(t.c.status, "Active") != "Merged"]
    used = False
    if q and q.strip():
        term = q.strip().upper()
        m = re.search(r"(UH\d{4,})", term)
        term = m.group(1) if m else term
        enc = (await ctx.conn.execute(sa.select(encounters.c.patient_id).where(sa.func.upper(encounters.c.code) == term))).first()
        conds.append(t.c.id == enc.patient_id if enc else sa.func.upper(t.c.code) == term)
        used = True
    if uhid and uhid.strip():
        if len(uhid.strip()) < 4:
            raise bad_request("Type at least 4 characters of the UHID")
        conds.append(t.c.code.ilike(f"%{uhid.strip()}%"))
        used = True
    if name and name.strip():
        if len(re.sub(r"[^A-Za-z]", "", name)) < 3:
            raise bad_request("Type at least 3 letters of the name")
        for part in name.strip().split()[:3]:
            conds.append(sa.func.lower(sa.func.concat(t.c.first_name, " ", t.c.last_name)).like(f"%{part.lower()}%"))
        used = True
    if aadhaar and aadhaar.strip():
        d = re.sub(r"\D", "", aadhaar)
        if len(d) < 4:
            raise bad_request("Enter the last 4 digits of the Aadhaar number")
        conds.append(t.c.id_number.like(f"%{d[-4:]}"))  # stored masked — only the last 4 digits are kept
        used = True
    if mobile and mobile.strip():
        d = re.sub(r"\D", "", mobile)
        if len(d) < 5:
            raise bad_request("Type at least 5 digits of the mobile number")
        conds.append(t.c.phone.like(f"%{d[-10:]}%"))
        used = True
    if not used:
        return {"rows": []}
    rows = (await ctx.conn.execute(sa.select(t).where(*conds).order_by(t.c.created_at.desc()).limit(max(1, min(limit, 50))))).all()
    recs = [rs.to_record(r) for r in rows]
    pids = [r["id"] for r in recs]
    open_adm: dict[str, dict] = {}
    if pids:
        e, b = encounters, beds
        w = records.alias("w")
        for a in (await ctx.conn.execute(
                sa.select(e.c.id, e.c.code, e.c.status, e.c.patient_id, b.c.bed_no, w.c.title.label("ward"))
                .select_from(e.outerjoin(b, b.c.id == e.c.bed_id).outerjoin(w, w.c.id == b.c.ward_id))
                .where(e.c.patient_id.in_(pids), e.c.type == "IPD", e.c.status.in_(OPEN), e.c.deleted_at.is_(None)))).all():
            open_adm[str(a.patient_id)] = {"id": str(a.id), "code": a.code, "status": a.status, "bed": a.bed_no, "ward": a.ward}
    out = []
    for r in recs:
        r["age"] = age_from(r.get("dob"))
        out.append({**_patient_brief(r), "admission": open_adm.get(r["id"])})
    return {"rows": out}


async def _ward_rows(ctx) -> list[dict]:
    ent = E(WARDS)
    t = ent.table
    rows = (await ctx.conn.execute(sa.select(t.c.id, t.c.code, t.c.title, t.c.status, t.c.data).where(*rs.base_where(ent)))).all()
    vacant = {str(w): n for w, n in (await ctx.conn.execute(
        sa.select(beds.c.ward_id, sa.func.count()).where(beds.c.deleted_at.is_(None), beds.c.status == "Vacant")
        .group_by(beds.c.ward_id))).all()}
    total = {str(w): n for w, n in (await ctx.conn.execute(
        sa.select(beds.c.ward_id, sa.func.count()).where(beds.c.deleted_at.is_(None)).group_by(beds.c.ward_id))).all()}
    out = []
    for r in rows:
        d = r.data or {}
        out.append({"id": str(r.id), "code": r.code, "name": r.title or r.code, "floor": d.get("floor"), "block": d.get("block"),
                    "ward_type": d.get("ward_type"), "department_id": d.get("department"), "status": r.status,
                    "vacant": vacant.get(str(r.id), 0), "beds": total.get(str(r.id), 0)})
    return sorted(out, key=lambda w: (str(w.get("floor") or ""), w["name"]))


@router.get("/admission-form")
async def admission_form(ctx: CtxDep):
    """Everything the New Admission screen needs in one call."""
    rs.ensure_access(ctx, E(ADM), "read")
    doc = employees
    dept = E(DEPTS)
    dept_rows =(await ctx.conn.execute(sa.select(dept.table.c.id, dept.table.c.code, dept.table.c.title, dept.table.c.data)
                                        .where(*rs.base_where(dept)).order_by(dept.table.c.title))).all()
    dept_names = {str(r.id): r.title for r in dept_rows}
    doctors = (await ctx.conn.execute(sa.select(doc.c.id, doc.c.name, doc.c.department_id, doc.c.designation, doc.c.code)
                                      .where(doc.c.deleted_at.is_(None), doc.c.category == "Doctor",
                                             sa.func.coalesce(doc.c.status, "Active") == "Active").order_by(doc.c.name))).all()
    payers = [{"id": p["id"], "name": p.get("name") or p.get("title"), "payer_type": p.get("payer_type"), "scheme": p.get("scheme"),
               "code": p.get("code")} for p in await find(ctx, PAYERS, status="Active", limit=500, sort="title")]
    packages = [{"id": p["id"], "name": p.get("name") or p.get("title"), "price": p.get("price"), "los_days": p.get("los_days"),
                 "room_category": p.get("room_category"), "specialty": p.get("specialty")}
                for p in await find(ctx, PACKAGES, status="Active", limit=500)]
    fm = E(ADM).field_map
    return {
        "wards": await _ward_rows(ctx),
        "departments": [{"id": str(r.id), "code": r.code, "name": r.title, "type": (r.data or {}).get("dept_type")} for r in dept_rows],
        "doctors": [{"id": str(r.id), "name": r.name, "code": r.code, "designation": r.designation,
                     "department_id": str(r.department_id) if r.department_id else None,
                     "department": dept_names.get(str(r.department_id)) if r.department_id else None} for r in doctors],
        "payers": payers, "packages": packages,
        "categories": [{"key": k, "label": v} for k, v in PAYER_LABEL.items()],
        "sponsor_types": {k: sorted(v) for k, v in SPONSOR_TYPES.items()},
        "admission_types": fm["admission_type"].options, "admission_sources": fm["admission_source"].options,
        "can_admit": _perm(ctx, "admit"),
        "can_view_money": _money_ok(ctx),
    }


@router.get("/wards/{ward_id}/beds")
async def ward_beds(ward_id: str, ctx: CtxDep, patient_id: str | None = None):
    """Beds a patient can be admitted to in this ward: vacant, or reserved for this patient."""
    rs.ensure_access(ctx, E(BED), "read")
    if not is_uuid(ward_id):
        raise not_found("Ward not found")
    b = beds
    rows = (await ctx.conn.execute(sa.select(b.c.id, b.c.code, b.c.bed_no, b.c.category, b.c.rate, b.c.status, b.c.data)
                                   .where(b.c.deleted_at.is_(None), b.c.ward_id == ward_id, b.c.status.in_(["Vacant", "Reserved"]))
                                   .order_by(b.c.bed_no))).all()
    out = []
    for r in rows:
        d = r.data or {}
        if r.status == "Reserved" and d.get("reserved_for") not in (None, patient_id):
            continue
        out.append({"id": str(r.id), "code": r.code, "bed_no": r.bed_no or r.code, "room": d.get("room_no"), "category": r.category,
                    "rate": to_json_value(r.rate), "status": r.status, "reserved_ref": d.get("reserved_ref")})
    return {"beds": sorted(out, key=lambda x: (len(str(x["bed_no"])), str(x["bed_no"])))}


class AdmitIn(BaseModel):
    patient_id: str
    payer: Literal["Cash", "Insurance", "Corporate", "Government Scheme"] = "Cash"
    sponsor_id: str | None = None
    authorisation_no: str | None = Field(default=None, max_length=60)
    approved_amount: float | None = Field(default=None, ge=0, le=100_000_000)
    doctor_id: str
    second_doctor_id: str | None = None
    department_id: str | None = None
    ward_id: str
    bed_id: str
    reason: str | None = Field(default=None, max_length=4000)
    diagnosis: str = Field(min_length=2, max_length=500)
    admission_type: str = "Elective"
    admission_source: str = "OPD"
    package_admission: bool = False
    package_id: str | None = None
    expected_los: int | None = Field(default=None, ge=0, le=365)
    attendant_name: str | None = Field(default=None, max_length=120)
    attendant_phone: str | None = Field(default=None, max_length=20)
    mlc: bool = False


async def _doctor(ctx, eid: str | None, field: str) -> dict | None:
    if not eid:
        return None
    row = (await ctx.conn.execute(sa.select(employees.c.id, employees.c.name, employees.c.category, employees.c.department_id)
                                  .where(employees.c.id == eid, employees.c.deleted_at.is_(None)))).first() if is_uuid(eid) else None
    if not row or row.category != "Doctor":
        raise bad_request("Choose a doctor from the list", {"fields": {field: "Not a doctor"}})
    return {"id": str(row.id), "name": row.name, "department_id": str(row.department_id) if row.department_id else None}


async def _sponsor(ctx, payer: str, sponsor_id: str | None) -> dict | None:
    if payer == "Cash":
        return None
    if not sponsor_id:
        raise bad_request(f"Choose the {PAYER_LABEL[payer].lower()} organisation / sponsor", {"fields": {"sponsor_id": "Required"}})
    sp = await get(ctx, PAYERS, sponsor_id)
    if not sp:
        raise bad_request("Sponsor not found", {"fields": {"sponsor_id": "Unknown sponsor"}})
    if sp.get("status") != "Active":
        raise conflict(f"{sp.get('name')} is {str(sp.get('status')).lower()} — cashless / credit is not available")
    ok = SPONSOR_TYPES.get(payer, set())
    if sp.get("payer_type") and sp["payer_type"] not in ok:
        raise bad_request(f"{sp.get('name')} is registered as {sp['payer_type']} — pick the matching payer category",
                          {"fields": {"sponsor_id": "Category mismatch"}})
    return sp


@router.post("/admissions", status_code=201)
async def admit(body: AdmitIn, ctx: CtxDep):
    _need(ctx, "admit", "admit patients")
    rs.ensure_access(ctx, E(ADM), "write")
    fm = E(ADM).field_map
    if body.admission_type not in (fm["admission_type"].options or []):
        raise bad_request("Unknown admission type", {"fields": {"admission_type": "Choose from the list"}})
    if body.admission_source not in (fm["admission_source"].options or []):
        raise bad_request("Unknown admission source", {"fields": {"admission_source": "Choose from the list"}})
    doc = await _doctor(ctx, body.doctor_id, "doctor_id")
    second = await _doctor(ctx, body.second_doctor_id, "second_doctor_id")
    if second and second["id"] == doc["id"]:
        raise bad_request("The second consultant must be a different doctor", {"fields": {"second_doctor_id": "Same as primary"}})
    dept = body.department_id or doc.get("department_id")
    if not dept or not await get(ctx, DEPTS, dept):
        raise bad_request("Choose the admitting department", {"fields": {"department_id": "Required"}})
    bed = await get(ctx, BED, body.bed_id)
    if not bed:
        raise bad_request("Choose a bed", {"fields": {"bed_id": "Required"}})
    if bed.get("ward_id") != body.ward_id:
        raise bad_request(f"Bed {bed.get('bed_no')} is not in the selected ward", {"fields": {"bed_id": "Different ward"}})
    sp = await _sponsor(ctx, body.payer, body.sponsor_id)
    pkg = None
    if body.package_admission:
        pkg = await get(ctx, PACKAGES, body.package_id) if body.package_id else None
        if not pkg or pkg.get("status") != "Active":
            raise bad_request("Choose an active package for a package admission", {"fields": {"package_id": "Required"}})
    p = await _patient_row(ctx, body.patient_id)
    entered = {"patient_id": body.patient_id, "doctor_id": doc["id"], "second_doctor": second["id"] if second else None,
               "department_id": dept, "bed_id": body.bed_id, "admission_type": body.admission_type, "admission_source": body.admission_source,
               "reason": (body.reason or "").strip() or None, "diagnosis": body.diagnosis.strip(), "expected_los": body.expected_los,
               "payer": body.payer, "sponsor": sp["id"] if sp else None, "package": pkg["id"] if pkg else None,
               "authorisation_no": (body.authorisation_no or "").strip() or None if sp else None,
               "approved_amount": body.approved_amount if sp else None, "attendant_name": body.attendant_name,
               "attendant_phone": body.attendant_phone, "mlc": body.mlc}
    if _check_screen is not None:  # field configuration (required / hidden / defaults) for the admission screen
        cleaned = await _check_screen(ctx, "ip.admission", entered)
        for k in ("patient_id", "doctor_id", "department_id", "bed_id", "payer", "sponsor", "diagnosis"):
            cleaned[k] = entered[k]  # system-required for a valid admission
        entered = cleaned
    payload = {**entered, "sponsor_name": sp.get("name") if sp else None, "package_admission": bool(pkg),
               "condition": "Not assessed", "admitted_by": _who(ctx), "status": "Admitted"}
    rec = await create(ctx, ADM, payload)
    ward = await get(ctx, WARDS, body.ward_id)
    where = f"{(ward or {}).get('name') or 'ward'} · bed {bed.get('bed_no')}"
    await notify_step(ctx, rec, "admitted", f"New admission {rec['code']}: {p.get('title')} ({p.get('code')}) → {where}. "
                                             "Accept the patient on the ward to start the nursing workflow.")
    return {"admission": rec, "patient": _patient_brief(p), "bed": {"id": bed["id"], "bed_no": bed.get("bed_no")},
            "ward": {"id": body.ward_id, "name": (ward or {}).get("name")}}


# ================================================================== admissions list

def _list_query():
    e, p, b, doc = encounters, patients, beds, employees
    w = records.alias("w")
    sp = records.alias("sp")
    q = (sa.select(e.c.id, e.c.code, e.c.status, e.c.patient_id, e.c.bed_id, e.c.started_at, e.c.ended_at, e.c.data,
                   p.c.code.label("uhid"), p.c.first_name, p.c.last_name, p.c.gender, p.c.dob, p.c.phone,
                   b.c.bed_no, b.c.ward_id, w.c.title.label("ward"), doc.c.name.label("doctor"), sp.c.title.label("sponsor_title"))
         .select_from(e.join(p, p.c.id == e.c.patient_id).outerjoin(b, b.c.id == e.c.bed_id).outerjoin(w, w.c.id == b.c.ward_id)
                      .outerjoin(doc, doc.c.id == e.c.doctor_id)
                      .outerjoin(sp, sa.cast(sp.c.id, sa.Text) == e.c.data["sponsor"].astext))
         .where(e.c.type == "IPD", e.c.deleted_at.is_(None)))
    return q, w


def _list_row(r) -> dict:
    d = r.data or {}
    return {
        "id": str(r.id), "code": r.code, "status": r.status, "started_at": _iso(r.started_at), "ended_at": _iso(r.ended_at),
        "patient": {"id": str(r.patient_id), "name": _name(r.first_name, r.last_name), "uhid": r.uhid, "gender": r.gender,
                    "age": age_from(r.dob), "phone": r.phone},
        "ward": {"id": str(r.ward_id) if r.ward_id else None, "name": r.ward}, "bed_no": r.bed_no, "doctor": r.doctor,
        "payer": d.get("payer") or "Cash", "payer_label": PAYER_LABEL.get(d.get("payer") or "Cash", d.get("payer")),
        "sponsor": {"id": d.get("sponsor"), "name": r.sponsor_title or d.get("sponsor_name")} if d.get("sponsor") else None,
        "condition": d.get("condition") or "Not assessed", "condition_at": d.get("condition_at"), "condition_by": d.get("condition_by"),
        "accepted_at": d.get("accepted_at"), "accepted_by": d.get("accepted_by"), "diagnosis": d.get("diagnosis"),
        "discharge_type": d.get("discharge_type"), "isolation": d.get("isolation"), "mlc": bool(d.get("mlc")),
        "los_days": _los_days(r.started_at, r.ended_at if r.status in ("Discharged", "Cancelled") else None),
    }


@router.get("/admissions")
async def admissions(ctx: CtxDep, q: str | None = None, status: str | None = None, ward: str | None = None,
                     sponsor: str | None = None, acceptance: Literal["accepted", "pending"] | None = None,
                     condition: str | None = None, date_from: str | None = None, date_to: str | None = None,
                     sort: str = "-started_at", page: int = 1, limit: int = 25):
    return await admissions_data(ctx, q=q, status=status, ward=ward, sponsor=sponsor, acceptance=acceptance, condition=condition,
                                 date_from=date_from, date_to=date_to, sort=sort, page=page, limit=limit)


async def admissions_data(ctx, *, q: str | None = None, status: str | None = None, ward: str | None = None, sponsor: str | None = None,
                          acceptance: str | None = None, condition: str | None = None, date_from: str | None = None,
                          date_to: str | None = None, sort: str = "-started_at", page: int = 1, limit: int = 25, cap: int = 200) -> dict:
    rs.ensure_access(ctx, E(ADM), "read")
    e, p, b = encounters, patients, beds
    base, w = _list_query()
    conds = []
    if q and q.strip():
        term = f"%{q.strip().lower()}%"
        digits = re.sub(r"\D", "", q)
        ors = [sa.func.lower(e.c.code).like(term), sa.func.lower(p.c.code).like(term),
               sa.func.lower(sa.func.concat(p.c.first_name, " ", p.c.last_name)).like(term)]
        if len(digits) >= 5:
            ors.append(p.c.phone.like(f"%{digits}%"))
        conds.append(sa.or_(*ors))
    if ward:
        if not is_uuid(ward):
            raise bad_request("Invalid ward")
        conds.append(b.c.ward_id == ward)
    if sponsor == "self":
        conds.append(sa.func.coalesce(e.c.data["sponsor"].astext, "") == "")
    elif sponsor:
        conds.append(e.c.data["sponsor"].astext == sponsor)
    if acceptance == "accepted":
        conds.append(e.c.data["accepted_at"].astext.isnot(None))
    elif acceptance == "pending":
        conds += [e.c.data["accepted_at"].astext.is_(None), e.c.status.in_(OPEN)]
    if condition:
        conds.append(sa.func.coalesce(e.c.data["condition"].astext, "Not assessed") == condition)
    if date_from:
        conds.append(e.c.started_at >= parse_dt(date_from))
    if date_to:
        conds.append(e.c.started_at < parse_dt(date_to) + timedelta(days=1))
    sub = base.where(*conds).subquery()  # KPI tiles count across statuses (the status filter applies to the rows)
    counts = {s: n for s, n in (await ctx.conn.execute(sa.select(sub.c.status, sa.func.count()).group_by(sub.c.status))).all()}
    crit = (await ctx.conn.execute(sa.select(sa.func.count()).select_from(sub).where(
        sub.c.status.in_(OPEN), sub.c.data["condition"].astext == "Critical"))).scalar_one()
    pending_accept = (await ctx.conn.execute(sa.select(sa.func.count()).select_from(sub).where(
        sub.c.status.in_(OPEN), sub.c.data["accepted_at"].astext.is_(None)))).scalar_one()
    if status:
        conds.append(e.c.status.in_([s for s in status.split(",") if s]))
    orders_by = {"started_at": e.c.started_at, "code": e.c.code, "status": e.c.status, "payer": e.c.data["payer"].astext,
                 "patient": p.c.first_name, "ward": w.c.title}
    key = sort.lstrip("-")
    col = orders_by.get(key, e.c.started_at)
    order = col.desc().nulls_last() if sort.startswith("-") else col.asc().nulls_last()
    limit = max(1, min(limit, cap))
    page = max(1, page)
    total = (await ctx.conn.execute(sa.select(sa.func.count()).select_from(base.where(*conds).subquery()))).scalar_one()
    rows = (await ctx.conn.execute(base.where(*conds).order_by(order, e.c.created_at.desc()).limit(limit).offset((page - 1) * limit))).all()
    return {
        "rows": [_list_row(r) for r in rows], "total": total, "page": page, "limit": limit,
        "kpis": {"total": sum(counts.values()), "admitted": counts.get("Admitted", 0), "discharge_initiated": counts.get("Discharge Initiated", 0),
                 "discharged": counts.get("Discharged", 0), "cancelled": counts.get("Cancelled", 0), "critical": crit,
                 "awaiting_acceptance": pending_accept},
        "filters": {"wards": [{"id": x["id"], "name": x["name"]} for x in await _ward_rows(ctx)],
                    "sponsors": [{"id": s["id"], "name": s.get("name") or s.get("title")} for s in await find(ctx, PAYERS, limit=500, sort="title")],
                    "conditions": CONDITIONS},
        "can": {"admit": _perm(ctx, "admit"), "accept": _perm(ctx, "accept"), "documents": _ent_ok(ctx, DOCS), "money": _money_ok(ctx)},
    }


# ================================================================== chart header (banner)

async def _bed_info(ctx, bed_id: str | None) -> dict | None:
    if not bed_id or not is_uuid(bed_id):
        return None
    b = beds
    w = records.alias("w")
    r = (await ctx.conn.execute(sa.select(b.c.id, b.c.bed_no, b.c.category, b.c.rate, b.c.ward_id, b.c.data, w.c.title.label("ward"),
                                          w.c.data.label("wdata")).select_from(b.outerjoin(w, w.c.id == b.c.ward_id))
                                .where(b.c.id == bed_id))).first()
    if not r:
        return None
    return {"id": str(r.id), "bed_no": r.bed_no, "category": r.category, "rate": to_json_value(r.rate), "room": (r.data or {}).get("room_no"),
            "ward_id": str(r.ward_id) if r.ward_id else None, "ward": r.ward, "floor": (r.wdata or {}).get("floor"),
            "block": (r.wdata or {}).get("block")}


async def allergy_view(ctx, p: dict) -> dict:
    review = p.get("allergy_review") or {}
    known = [a["allergen"] for a in await patient_allergies(ctx, p["id"], p)]
    types = list(review.get("types") or [])
    if review.get("nka") and not known:
        status = "nka"
    elif known or types or review.get("substance"):
        status = "present"
    else:
        status = "unknown"
    return {"status": status, "nka": bool(review.get("nka")) and not known, "types": types, "substance": review.get("substance"),
            "list": known, "by": review.get("by"), "at": review.get("at")}


def reactive_view(adm: dict) -> dict:
    r = adm.get("reactive") or {}
    conds = list(r.get("conditions") or [])
    status = "non_reactive" if r.get("non_reactive") else "reactive" if conds or r.get("other") else "unknown"
    return {"status": status, "conditions": conds, "other": r.get("other"), "by": r.get("by"), "at": r.get("at"),
            "isolation": adm.get("isolation") or "None"}


async def _latest_vitals(ctx, pid: str) -> dict | None:
    t = E(VITALS).table
    rows = await _rows(ctx, VITALS, t.c.patient_id == pid, sa.func.coalesce(t.c.status, "") != "Entered in Error",
                       order=[_dcol(VITALS, "recorded_at").desc().nulls_last(), t.c.created_at.desc()], limit=1)
    return _vital_point(rows[0]) if rows else None


async def _count(ctx, fqn: str, *conds) -> int:
    if not has(fqn):
        return 0
    ent = E(fqn)
    return (await ctx.conn.execute(sa.select(sa.func.count()).select_from(ent.table).where(*rs.base_where(ent), *conds))).scalar_one()


async def _tab_counts(ctx, adm: dict) -> dict:
    aid, pid = adm["id"], adm["patient_id"]
    o = orders
    oc = {(c, s): n for c, s, n in (await ctx.conn.execute(sa.select(o.c.category, o.c.status, sa.func.count()).where(
        o.c.encounter_id == aid, o.c.deleted_at.is_(None)).group_by(o.c.category, o.c.status))).all()}
    inv = sum(n for (c, s), n in oc.items() if c in ("LAB", "RAD") and s != "Cancelled")
    active = sum(n for (c, s), n in oc.items() if s in ("Ordered", "Collected", "In Progress"))
    ti = E(INDENT).table if has(INDENT) else None
    raised = await _count(ctx, INDENT, ti.c.encounter_id == aid, ti.c.status == "Raised") if ti is not None else 0
    tt = E(TASKS).table if has(TASKS) else None
    tasks = await _count(ctx, TASKS, tt.c.patient_id == pid, tt.c.status == "Pending") if tt is not None else 0
    ins = E(INSTR).table
    instr = await _count(ctx, INSTR, ins.c.encounter_id == aid, ins.c.status == "Open")
    docs = await _count(ctx, DOCS, E(DOCS).table.c.encounter_id == aid) if has(DOCS) else 0
    cons = await _count(ctx, CONSENT, E(CONSENT).table.c.encounter_id == aid) if has(CONSENT) else 0
    ot = await _count(ctx, OT_BOOKINGS, E(OT_BOOKINGS).table.c.patient_id == pid,
                      E(OT_BOOKINGS).table.c.status.notin_(["Cancelled", "Shifted Out"])) if has(OT_BOOKINGS) else 0
    us = E(USAGE).table
    eq = await _count(ctx, USAGE, us.c.encounter_id == aid, us.c.status == "Connected")
    return {"documents": docs, "consents": cons, "investigations": inv, "order_desk": active + raised + eq, "nursing": tasks,
            "doctor": instr, "ot": ot, "equipment": eq, "indents": raised}


def _perms(ctx) -> dict[str, Any]:
    out: dict[str, Any] = {k: _perm(ctx, k) for k in PERMS}
    clinical = _can(ctx, *CLINICAL_VIEW) and _ent_ok(ctx, VITALS)
    desk = any(out[k] for k in ("order", "indent", "charge", "collect", "equipment", "pharmacy_return", "money"))
    out["tabs"] = {
        "overview": True, "forms": clinical, "documents": _ent_ok(ctx, DOCS), "consent": _ent_ok(ctx, CONSENT),
        "cath": clinical, "dietary": clinical, "doctor": clinical and _ent_ok(ctx, NOTES), "hic": clinical, "monitoring": clinical,
        "nursing": clinical and (_ent_ok(ctx, MAR) or _ent_ok(ctx, TASKS)), "order_desk": desk,
        "order_medication": out["indent"] or (_ent_ok(ctx, INDENT) and _can(ctx, "pharmacist", "management", "doctor", "nurse")),
        "investigations": _ent_ok(ctx, ORD) or out["collect"], "ot": clinical and _ent_ok(ctx, OT_BOOKINGS),
    }
    out["clearance_roles"] = {k: sorted(v) for k, v in CLEARANCE_ROLES.items()}
    return out


@router.get("/admissions/{aid}")
async def chart_header(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(ADM), "read")
    adm = await _admission(ctx, aid)
    p = await _patient_row(ctx, adm["patient_id"])
    await audit(ctx, "ipchart.view", module="ipd", entity="admissions", record=adm)
    names = await _emp_names(ctx, [adm.get("doctor_id"), adm.get("second_doctor")])
    doc_dept = {}
    if adm.get("doctor_id") and is_uuid(adm["doctor_id"]):
        r = (await ctx.conn.execute(sa.select(employees.c.designation, employees.c.department_id).where(employees.c.id == adm["doctor_id"]))).first()
        if r:
            doc_dept = {"designation": r.designation, "department_id": str(r.department_id) if r.department_id else None}
    dept = await get(ctx, DEPTS, adm.get("department_id")) if adm.get("department_id") else None
    sponsor = await get(ctx, PAYERS, adm.get("sponsor")) if adm.get("sponsor") else None
    pkg = await get(ctx, PACKAGES, adm.get("package")) if adm.get("package") else None
    track = await tracker_state(ctx, adm)
    money_ok = _money_ok(ctx)
    return {
        "admission": {
            **{k: adm.get(k) for k in ("id", "code", "status", "started_at", "ended_at", "diagnosis", "reason", "admission_type",
                                      "admission_source", "payer", "authorisation_no", "package_admission", "expected_los",
                                      "attendant_name", "attendant_phone", "mlc", "isolation", "discharge_type", "discharge_initiated_at",
                                      "discharge_initiated_by", "accepted_at", "accepted_by", "condition_at", "condition_by", "admitted_by",
                                      "created_at")},
            "condition": adm.get("condition") or "Not assessed",
            "condition_log": list(reversed((adm.get("condition_log") or [])[-20:])),
            "sponsor_history": list(reversed((adm.get("sponsor_history") or [])[-20:])),
            "payer_label": PAYER_LABEL.get(adm.get("payer") or "Cash", adm.get("payer")),
            "sponsor": {"id": sponsor["id"], "name": sponsor.get("name"), "payer_type": sponsor.get("payer_type")} if sponsor else None,
            "package": {"id": pkg["id"], "name": pkg.get("name"), "price": pkg.get("price") if money_ok else None} if pkg else None,
            "approved_amount": adm.get("approved_amount") if money_ok else None,
            "los_days": _los_days(adm.get("started_at"), adm.get("ended_at") if adm.get("status") in ("Discharged", "Cancelled") else None),
        },
        "patient": _patient_brief(p),
        "bed": await _bed_info(ctx, adm.get("bed_id")),
        "department": {"id": dept["id"], "name": dept.get("name") or dept.get("title")} if dept else None,
        "doctor": {"id": adm.get("doctor_id"), "name": names.get(adm.get("doctor_id") or ""), **doc_dept} if adm.get("doctor_id") else None,
        "second_doctor": {"id": adm["second_doctor"], "name": names.get(adm["second_doctor"])} if adm.get("second_doctor") else None,
        "allergies": await allergy_view(ctx, p),
        "reactive": reactive_view(adm),
        "last_obs": await _latest_vitals(ctx, adm["patient_id"]),
        "counts": await _tab_counts(ctx, adm),
        "tracker": {"done": track["done"], "total": len(STEPS), "next": track["next"], "steps": {s["key"]: s["state"] for s in track["steps"]}},
        "permissions": _perms(ctx),
        "options": {"conditions": CONDITIONS, "allergy_types": ALLERGY_TYPES, "reactive": REACTIVE, "shifts": SHIFTS,
                    "discharge_types": E(ADM).field_map["discharge_type"].options, "isolation": ISOLATION},
    }


class AcceptIn(BaseModel):
    remarks: str | None = Field(default=None, max_length=300)


@router.post("/admissions/{aid}/accept")
async def accept_patient(aid: str, body: AcceptIn, ctx: CtxDep):
    """Ward acceptance: the receiving nurse confirms the patient has arrived on the ward."""
    _need(ctx, "accept", "accept patients on the ward")
    adm = await _admission(ctx, aid)
    _require_open(adm, "accept the patient")
    if adm.get("accepted_at"):
        raise conflict(f"Already accepted by {adm.get('accepted_by')} at {parse_dt(adm['accepted_at']).astimezone(settings.tz):%d %b %H:%M}")
    rec = await update(ctx, ADM, aid, {"accepted_at": _stamp(), "accepted_by": _who(ctx), "accepted_by_user": ctx.user_id,
                                       "acceptance_remarks": (body.remarks or "").strip() or None})
    await audit(ctx, "ward.accept", module="ipd", entity="admissions", record=rec, note=body.remarks)
    return {"accepted_at": rec.get("accepted_at"), "accepted_by": rec.get("accepted_by")}


class ConditionIn(BaseModel):
    condition: str
    note: str | None = Field(default=None, max_length=300)


@router.post("/admissions/{aid}/condition")
async def set_condition(aid: str, body: ConditionIn, ctx: CtxDep):
    _need(ctx, "condition", "set the clinical condition")
    if body.condition not in CONDITIONS:
        raise bad_request("Choose a condition from the list")
    adm = await _admission(ctx, aid)
    _require_open(adm, "change the condition")
    prev = adm.get("condition") or "Not assessed"
    if prev == body.condition and not body.note:
        return {"condition": prev, "condition_at": adm.get("condition_at"), "condition_by": adm.get("condition_by")}
    at = _stamp()
    entry = {"condition": body.condition, "previous": prev, "by": _who(ctx), "role": ctx.role, "at": at, "note": (body.note or "").strip() or None}
    log_ = [*(adm.get("condition_log") or []), entry][-100:]
    rec = await update(ctx, ADM, aid, {"condition": body.condition, "condition_at": at, "condition_by": _who(ctx), "condition_log": log_})
    if body.condition in ("Critical", "Serious") and prev not in ("Critical", "Serious"):
        bed = await _bed_info(ctx, adm.get("bed_id"))
        p = await _patient_row(ctx, adm["patient_id"])
        live(ctx, "alerts", {"type": "condition", "roles": ["doctor", "nurse"], "patient_id": adm["patient_id"], "admission_id": aid,
                             "message": f"{p.get('title')} ({adm.get('code')}) marked {body.condition.upper()} by {_who(ctx)}"
                                        f"{' · ' + str(bed['ward']) + ' bed ' + str(bed['bed_no']) if bed else ''}"})
    return {"condition": rec.get("condition"), "condition_at": rec.get("condition_at"), "condition_by": rec.get("condition_by"), "log_entry": entry}


class AllergyIn(BaseModel):
    nka: bool = False
    types: list[str] = Field(default_factory=list, max_length=9)
    substance: str | None = Field(default=None, max_length=500)


@router.put("/admissions/{aid}/allergies")
async def save_allergies(aid: str, body: AllergyIn, ctx: CtxDep):
    """Patient-level allergy status. Named substances join the patient's known-allergy list, which the
    pharmacy safety screen checks on every prescription, indent and dispense."""
    _need(ctx, "clinical", "record allergies")
    adm = await _admission(ctx, aid)
    bad = [t for t in body.types if t not in ALLERGY_TYPES]
    if bad:
        raise bad_request(f"Unknown allergy type: {', '.join(bad)}")
    substance = (body.substance or "").strip() or None
    if body.nka and (body.types or substance):
        raise bad_request("Either record the allergies or confirm 'No known allergies' — not both")
    if not body.nka and not body.types and not substance:
        raise bad_request("Tick an allergy type, name the substance, or confirm 'No known allergies'")
    p = await _patient_row(ctx, adm["patient_id"])
    if body.nka:
        known = [a["allergen"] for a in await patient_allergies(ctx, p["id"], p) if a.get("source") == "EMR"]
        if known:
            raise conflict(f"The EMR has active allergies ({', '.join(known[:4])}) — mark them inactive there before confirming none")
    review = {"nka": body.nka, "types": [t for t in ALLERGY_TYPES if t in body.types], "substance": substance,
              "by": _who(ctx), "role": ctx.role, "at": _stamp(), "admission": adm.get("code")}
    tags = [] if body.nka else list(p.get("allergies") or [])
    for s in re.split(r"[,;/\n]+", substance or ""):
        s = s.strip()
        if s and s.lower() not in {t.lower() for t in tags}:
            tags.append(s[:80])
    await update(ctx, PATIENTS, p["id"], {"allergy_review": review, "allergies": tags or None})
    await audit(ctx, "allergies.review", module="ipd", entity="admissions", record=adm, after=review)
    p = await _patient_row(ctx, adm["patient_id"])
    return await allergy_view(ctx, p)


class ReactiveIn(BaseModel):
    non_reactive: bool = False
    conditions: list[str] = Field(default_factory=list, max_length=10)
    other: str | None = Field(default=None, max_length=300)


@router.put("/admissions/{aid}/reactive")
async def save_reactive(aid: str, body: ReactiveIn, ctx: CtxDep):
    """Reactive / infection-precaution status for this stay. Airborne, droplet, contact or protective
    precautions are raised on the admission (and shown on the bed board); they are never lowered here."""
    _need(ctx, "clinical", "record the reactive status")
    adm = await _admission(ctx, aid)
    _require_open(adm, "change the reactive status")
    bad = [c for c in body.conditions if c not in REACTIVE]
    if bad:
        raise bad_request(f"Unknown condition: {', '.join(bad)}")
    other = (body.other or "").strip() or None
    if body.non_reactive and (body.conditions or other):
        raise bad_request("Either tick the reactive conditions or confirm 'Non reactive' — not both")
    if not body.non_reactive and not body.conditions and not other:
        raise bad_request("Tick a condition or confirm 'Non reactive'")
    review = {"non_reactive": body.non_reactive, "conditions": [c for c in REACTIVE if c in body.conditions], "other": other,
              "by": _who(ctx), "role": ctx.role, "at": _stamp()}
    data: dict[str, Any] = {"reactive": review}
    cur = adm.get("isolation") or "None"
    need = max([ISOLATION_FOR[c] for c in review["conditions"] if c in ISOLATION_FOR] or ["None"], key=lambda x: ISOLATION_RANK.get(x, 0))
    raised = None
    if ISOLATION_RANK.get(need, 0) > ISOLATION_RANK.get(cur, 0):
        data["isolation"] = raised = need
    rec = await update(ctx, ADM, aid, data)
    if raised:
        bed = await _bed_info(ctx, adm.get("bed_id"))
        live(ctx, "alerts", {"type": "isolation", "roles": ["nurse", "support"], "admission_id": aid,
                             "message": f"{raised} precautions for {adm.get('code')}"
                                        f"{' · ' + str(bed['ward']) + ' bed ' + str(bed['bed_no']) if bed else ''}"})
    return {**reactive_view(rec), "isolation_raised": raised}


class SponsorIn(BaseModel):
    payer: Literal["Cash", "Insurance", "Corporate", "Government Scheme"]
    sponsor_id: str | None = None
    authorisation_no: str | None = Field(default=None, max_length=60)
    approved_amount: float | None = Field(default=None, ge=0, le=100_000_000)
    reason: str = Field(min_length=3, max_length=500)


@router.post("/admissions/{aid}/sponsor")
async def change_sponsor(aid: str, body: SponsorIn, ctx: CtxDep):
    """Record a payer / sponsor change for the stay (re-pricing of the bill is done in IP billing)."""
    _need(ctx, "sponsor", "change the sponsor")
    adm = await _admission(ctx, aid)
    _require_open(adm, "change the sponsor")
    sp = await _sponsor(ctx, body.payer, body.sponsor_id)
    before = {"payer": adm.get("payer") or "Cash", "sponsor": adm.get("sponsor"), "sponsor_name": adm.get("sponsor_name"),
              "authorisation_no": adm.get("authorisation_no"), "approved_amount": adm.get("approved_amount")}
    after = {"payer": body.payer, "sponsor": sp["id"] if sp else None, "sponsor_name": sp.get("name") if sp else None,
             "authorisation_no": (body.authorisation_no or "").strip() or None if sp else None,
             "approved_amount": body.approved_amount if sp else None}
    if before == after:
        raise bad_request("Nothing changed")
    hist = [*(adm.get("sponsor_history") or []), {"from": before, "to": after, "by": _who(ctx), "role": ctx.role, "at": _stamp(),
                                                  "reason": body.reason.strip()}][-50:]
    rec = await update(ctx, ADM, aid, {**after, "sponsor_history": hist, "repricing_pending": True})
    await audit(ctx, "admission.sponsor_change", module="ipd", entity="admissions", record=rec, before=before, after=after, note=body.reason)
    await emit("admission.sponsor_changed", ctx, {"admission": rec, "before": before, "after": after})
    return {"payer": rec.get("payer"), "sponsor": after, "history": list(reversed(hist))}


class ChangePatientIn(BaseModel):
    patient_id: str
    reason: str = Field(min_length=5, max_length=500)


@router.post("/admissions/{aid}/change-patient")
async def change_patient(aid: str, body: ChangePatientIn, ctx: CtxDep):
    """Fix an admission opened against the wrong UHID: the admission and everything recorded under it
    (orders, charges, vitals, notes, indents …) move to the correct patient. Audited; front office / admin."""
    _need(ctx, "change_patient", "change the patient on an admission")
    adm = await _admission(ctx, aid)
    _require_open(adm, "change the patient")
    old = adm["patient_id"]
    if body.patient_id == old:
        raise bad_request("This admission already belongs to that patient")
    new = await _patient_row(ctx, body.patient_id)
    if new.get("status") in ("Merged", "Deceased"):
        raise conflict(f"That patient record is marked {str(new.get('status')).lower()}")
    other = await first(ctx, ADM, patient_id=body.patient_id, status=list(OPEN))
    if other:
        raise conflict(f"{new.get('title')} already has an open admission ({other['code']})")
    billed = (await ctx.conn.execute(sa.select(sa.func.count()).where(
        invoices.c.encounter_id == aid, invoices.c.deleted_at.is_(None), invoices.c.status != "Cancelled"))).scalar_one()
    if billed:
        raise conflict("Bills have already been raised on this admission — cancel them in billing before re-linking the patient")
    oldp = await _patient_row(ctx, old)
    moved: dict[str, int] = {}
    for name, t in ENTITY_TABLES.items():
        if name in ("patients", "encounters") or "encounter_id" not in t.c or "patient_id" not in t.c:
            continue
        res = await ctx.conn.execute(t.update().where(t.c.encounter_id == aid, t.c.patient_id == old)
                                     .values(patient_id=body.patient_id, updated_at=now(), updated_by=ctx.user_id))
        if res.rowcount:
            moved[name] = res.rowcount
    log_ = [*(adm.get("patient_change_log") or []), {"from": {"id": old, "uhid": oldp.get("code"), "name": oldp.get("title")},
                                                    "to": {"id": new["id"], "uhid": new.get("code"), "name": new.get("title")},
                                                    "by": _who(ctx), "at": _stamp(), "reason": body.reason.strip()}]
    rec = await update(ctx, ADM, aid, {"patient_id": body.patient_id, "patient_change_log": log_})
    bed = await get(ctx, BED, adm.get("bed_id")) if adm.get("bed_id") else None
    if bed and bed.get("current_encounter") == aid:
        await update(ctx, BED, bed["id"], {"current_patient": body.patient_id})
    await audit(ctx, "admission.change_patient", module="ipd", entity="admissions", record=rec,
                before={"patient_id": old, "uhid": oldp.get("code")}, after={"patient_id": new["id"], "uhid": new.get("code"), "moved": moved},
                note=body.reason)
    return {"admission": {"id": aid, "code": rec.get("code")}, "patient": _patient_brief(new), "moved": moved,
            "note": "Advance deposits stay with the original UHID — refund or transfer them in billing if needed."}


class InitiateIn(BaseModel):
    discharge_type: str
    expected_at: str | None = None
    remarks: str | None = Field(default=None, max_length=500)


@router.post("/admissions/{aid}/initiate-discharge")
async def initiate_discharge(aid: str, body: InitiateIn, ctx: CtxDep):
    _need(ctx, "discharge", "initiate discharge")
    adm = await _admission(ctx, aid)
    if adm.get("status") != "Admitted":
        raise conflict(f"Admission {adm.get('code')} is {adm.get('status')}")
    if body.discharge_type not in (E(ADM).field_map["discharge_type"].options or []):
        raise bad_request("Choose the discharge type", {"fields": {"discharge_type": "Required"}})
    exp = parse_dt(body.expected_at) if body.expected_at else None
    if exp and exp < now() - timedelta(minutes=5):
        raise bad_request("Expected discharge time is in the past", {"fields": {"expected_at": "Past time"}})
    await transition(ctx, ADM, aid, "Discharge Initiated", note=body.remarks, values={"discharge_type": body.discharge_type})
    if exp or body.remarks:
        chk = await first(ctx, CHK, encounter_id=aid, sort="-created_at")
        if chk and chk.get("status") == "Open":
            await update(ctx, CHK, chk["id"], {k: v for k, v in {"expected_discharge_at": exp.isoformat() if exp else None,
                                                                 "remarks": body.remarks}.items() if v})
    return await tracker(aid, ctx)


class RevokeIn(BaseModel):
    reason: str = Field(min_length=3, max_length=500)


@router.post("/admissions/{aid}/revoke-discharge")
async def revoke_discharge(aid: str, body: RevokeIn, ctx: CtxDep):
    _need(ctx, "revoke_discharge", "revoke a discharge")
    adm = await _admission(ctx, aid)
    if adm.get("status") != "Discharge Initiated":
        raise conflict("Discharge has not been initiated")
    await transition(ctx, ADM, aid, "Admitted", note=body.reason)
    return await tracker(aid, ctx)


# ================================================================== vitals & observations

RANGE_HOURS = {"24h": 24, "3d": 72, "7d": 168, "30d": 720}


@router.get("/admissions/{aid}/vitals")
async def vitals(aid: str, ctx: CtxDep, window: Literal["24h", "3d", "7d", "30d"] = "24h"):
    rs.ensure_access(ctx, E(VITALS), "read")
    adm = await _admission(ctx, aid)
    t = E(VITALS).table
    at = _dcol(VITALS, "recorded_at")
    series = await _rows(ctx, VITALS, t.c.patient_id == adm["patient_id"], sa.func.coalesce(t.c.status, "") != "Entered in Error",
                         at >= now() - timedelta(hours=RANGE_HOURS[window]), order=[at.asc(), t.c.created_at.asc()], limit=2000)
    latest = await _latest_vitals(ctx, adm["patient_id"])
    names = await _emp_names(ctx, [latest.get("recorded_by")] if latest else [])
    if latest:
        latest["recorded_by_name"] = names.get(latest.get("recorded_by") or "")
    return {"window": window, "series": [_vital_point(v) for v in series], "latest": latest,
            "bands": {k: {"critical_low": v[0], "low": v[1], "high": v[2], "critical_high": v[3]} for k, v in RANGES.items()}}


class ObsIn(BaseModel):
    bp_systolic: float | None = Field(default=None, ge=40, le=300)
    bp_diastolic: float | None = Field(default=None, ge=20, le=200)
    pulse: float | None = Field(default=None, ge=20, le=250)
    resp_rate: float | None = Field(default=None, ge=2, le=80)
    spo2: float | None = Field(default=None, ge=40, le=100)
    temperature_c: float | None = Field(default=None, ge=32.3, le=43)
    blood_sugar: float | None = Field(default=None, ge=10, le=1000)
    pain_score: float | None = Field(default=None, ge=0, le=10)
    gcs: float | None = Field(default=None, ge=3, le=15)
    weight: float | None = Field(default=None, ge=0.3, le=400)
    height: float | None = Field(default=None, ge=20, le=250)
    o2_support: str | None = None
    taken_at: str | None = None
    remarks: str | None = Field(default=None, max_length=500)


@router.post("/admissions/{aid}/observations", status_code=201)
async def record_observations(aid: str, body: ObsIn, ctx: CtxDep):
    """'Chart it': records an emr.vitals entry (NEWS2 / BMI computed and escalated by clinical.py) and returns
    the parameters outside the normal adult range so the nurse sees them immediately."""
    _need(ctx, "observe", "record observations")
    adm = await _admission(ctx, aid)
    _require_open(adm, "record observations")
    vals = body.model_dump(exclude={"taken_at", "remarks", "o2_support"})
    if all(v is None for v in vals.values()):
        raise bad_request("Enter at least one observation")
    if (body.bp_systolic is None) != (body.bp_diastolic is None):
        raise bad_request("Enter both systolic and diastolic blood pressure", {"fields": {"bp_diastolic": "Both values needed"}})
    if body.bp_systolic is not None and body.bp_diastolic is not None and body.bp_diastolic >= body.bp_systolic:
        raise bad_request("Diastolic must be lower than systolic", {"fields": {"bp_diastolic": "Check the reading"}})
    taken = parse_dt(body.taken_at) if body.taken_at else now()
    if taken > now() + timedelta(minutes=5):
        raise bad_request("'Taken at' cannot be in the future", {"fields": {"taken_at": "Future time"}})
    started = parse_dt(adm.get("started_at"))
    if started and taken < started - timedelta(hours=12):
        raise bad_request("'Taken at' is before this admission", {"fields": {"taken_at": "Before admission"}})
    o2 = E(VITALS).field_map["o2_support"].options or []
    if body.o2_support and body.o2_support not in o2:
        raise bad_request("Unknown oxygen support option")
    payload = {
        "patient_id": adm["patient_id"], "encounter_id": aid, "recorded_at": taken.isoformat(timespec="minutes"),
        "bp_systolic": body.bp_systolic, "bp_diastolic": body.bp_diastolic, "pulse": body.pulse, "resp_rate": body.resp_rate,
        "spo2": body.spo2, "temperature": round(body.temperature_c * 9 / 5 + 32, 1) if body.temperature_c is not None else None,
        "blood_sugar": body.blood_sugar, "pain_score": body.pain_score, "gcs": body.gcs, "weight": body.weight, "height": body.height,
        "o2_support": body.o2_support or "Room air", "recorded_by": _emp(ctx), "remarks": body.remarks, "status": "Recorded",
        "source": "ipchart",
    }
    rec = await create(ctx, VITALS, {k: v for k, v in payload.items() if v is not None})
    point = _vital_point(rec)
    flags = _out_of_range(point)
    return {"vitals": point, "out_of_range": flags, "news2": rec.get("news2"), "news2_risk": rec.get("news2_risk")}


# ================================================================== overview widgets, visit tracker, timeline

def _stage(status: str | None) -> str:
    return {"Resulted": "Reported", "Verified": "Reported", "Cancelled": "Cancelled", "Collected": "Sample collected",
            "In Progress": "In progress"}.get(status or "", "Ordered")


def _order_q():
    o, s = orders, services
    return (sa.select(o.c.id, o.c.code, o.c.status, o.c.category, o.c.priority, o.c.ordered_at, o.c.ordered_by, o.c.verified_at,
                      o.c.resulted_at, o.c.collected_at, o.c.critical, o.c.data, o.c.service_id, s.c.code.label("svc_code"),
                      s.c.title.label("svc_name"), s.c.rate.label("svc_rate"))
            .select_from(o.outerjoin(s, s.c.id == o.c.service_id)).where(o.c.deleted_at.is_(None)))


async def _order_people(ctx, rows) -> tuple[dict[str, str], dict[str, str]]:
    users_ = await _user_names(ctx, [r.ordered_by for r in rows])
    docs = await _emp_names(ctx, [(r.data or {}).get("ordering_doctor") for r in rows])
    return users_, docs


FORM_ENTRIES = "emr.clinical_forms"  # clinical forms engine (forms.py)
INITIAL_FORMS = ("nursing.initial_assessment", "nursing.paediatric_assessment")
REASSESS_FORMS = ("nursing.morse", "nursing.braden", "nursing.mews", "nursing.pain", "nursing.gcs", "nursing.rass")
VISIT_FORMS = ("doctor.consultant_visit", "doctor.progress_soap", "doctor.admission_hp")


async def _form_entries(ctx, aid: str, keys: tuple[str, ...]) -> list[dict]:
    if not has(FORM_ENTRIES):
        return []
    t = E(FORM_ENTRIES).table
    return await _rows(ctx, FORM_ENTRIES, t.c.encounter_id == aid, t.c.status.in_(["Signed", "Amended"]),
                       t.c.data["form_key"].astext.in_(list(keys)), order=_dcol(FORM_ENTRIES, "noted_at").desc().nulls_last(), limit=20)


async def visit_tracker(ctx, adm: dict) -> dict:
    """Accept patient → initial nursing assessment → (optional) re-assessment in the last 24 h → consultant visit.
    Each step is satisfied by the structured forms (forms.py) or the legacy nursing assessments / progress notes."""
    aid, pid = adm["id"], adm["patient_id"]
    started = parse_dt(adm.get("started_at")) or now()
    items: list[dict] = [{"key": "accept", "label": "Accept patient", "optional": False, "done": bool(adm.get("accepted_at")),
                          "at": adm.get("accepted_at"), "by": adm.get("accepted_by"), "go": {"tab": "overview", "action": "accept"}}]
    init = rea = None
    if has(ASSESS):
        t = E(ASSESS).table
        rows = await _rows(ctx, ASSESS, t.c.patient_id == pid, _dcol(ASSESS, "assessed_at") >= started - timedelta(hours=6),
                           order=_dcol(ASSESS, "assessed_at").asc(), limit=200)
        init = next((r for r in rows if r.get("assessment_type") == "Admission"), None)
        later = [r for r in rows if r.get("assessment_type") != "Admission"
                 and parse_dt(r.get("assessed_at")) and parse_dt(r["assessed_at"]) >= now() - timedelta(hours=24)]
        rea = later[-1] if later else None
    f_init = next(iter(await _form_entries(ctx, aid, INITIAL_FORMS)), None)
    f_rea = next((f for f in await _form_entries(ctx, aid, REASSESS_FORMS)
                  if parse_dt(f.get("noted_at")) and parse_dt(f["noted_at"]) >= now() - timedelta(hours=24)), None)
    f_visit = next(iter(await _form_entries(ctx, aid, VISIT_FORMS)), None)
    names = await _emp_names(ctx, [(init or {}).get("assessed_by"), (rea or {}).get("assessed_by"), (f_init or {}).get("author"),
                                   (f_rea or {}).get("author"), (f_visit or {}).get("author")])

    def pick(form: dict | None, legacy: dict | None) -> tuple[Any, Any]:
        if form:
            return form.get("noted_at"), names.get(form.get("author") or "") or form.get("signed_by_name")
        if legacy:
            return legacy.get("assessed_at"), names.get(legacy.get("assessed_by") or "")
        return None, None

    at, by = pick(f_init, init)
    items.append({"key": "initial", "label": "Initial assessment", "optional": False, "done": bool(f_init or init), "at": at, "by": by,
                  "go": {"tab": "nursing", "sub": "nursing.initial_assessment"}})
    at, by = pick(f_rea, rea)
    items.append({"key": "reassessment", "label": "Adult re-assessment", "optional": True, "done": bool(f_rea or rea), "at": at, "by": by,
                  "go": {"tab": "nursing", "sub": "nursing.morse"}})
    visit = None
    nt = E(NOTES).table
    for n in await _rows(ctx, NOTES, nt.c.encounter_id == aid, order=_dcol(NOTES, "noted_at").desc().nulls_last(), limit=50):
        if n.get("note_type") in ("Progress Note", "Ward Round", "Consult Note", "History & Physical") and n.get("status") != "Draft":
            visit = n
            break
    seen = None
    if not visit and not f_visit and has(CONSULTS):
        ct = E(CONSULTS).table
        seen = next(iter(await _rows(ctx, CONSULTS, ct.c.encounter_id == aid, ct.c.status == "Seen", limit=1)), None)
    vn = await _emp_names(ctx, [(visit or {}).get("author")])
    v_at = (visit or {}).get("noted_at") or (f_visit or {}).get("noted_at") or (seen or {}).get("seen_at")
    v_by = vn.get((visit or {}).get("author") or "") or names.get((f_visit or {}).get("author") or "") or ("Cross-consult" if seen else None)
    items.append({"key": "consultant_visit", "label": "Consultant visit", "optional": False, "done": bool(visit or f_visit or seen),
                  "at": v_at, "by": v_by, "go": {"tab": "doctor", "sub": "progress-notes"}})
    req = [i for i in items if not i["optional"]]
    return {"items": items, "done": sum(1 for i in items if i["done"]), "total": len(items),
            "pending": sum(1 for i in req if not i["done"])}


@router.get("/admissions/{aid}/overview")
async def overview(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(ADM), "read")
    adm = await _admission(ctx, aid)
    pid = adm["patient_id"]
    o = orders
    rows = (await ctx.conn.execute(_order_q().where(o.c.encounter_id == aid).order_by(o.c.ordered_at.desc().nulls_last()).limit(8))).all()
    users_, docs = await _order_people(ctx, rows)
    inv = [{"id": str(r.id), "code": r.code, "test": r.svc_name or r.svc_code, "category": r.category, "status": r.status,
            "stage": _stage(r.status), "priority": r.priority or "Routine", "ordered_at": _iso(r.ordered_at), "critical": bool(r.critical),
            "ordered_by": docs.get((r.data or {}).get("ordering_doctor") or "") or users_.get(r.ordered_by or "")} for r in rows]

    meds: list[dict] = []
    if has(INDENT):
        ti = E(INDENT).table
        ind = await _rows(ctx, INDENT, ti.c.encounter_id == aid, order=ti.c.created_at.desc(), limit=6)
        its = await items_by_id(ctx, [ln.get("item") for r in ind for ln in r.get("items") or []])
        for r in ind:
            for ln in r.get("items") or []:
                meds.append({"name": item_label(its.get(ln.get("item"))), "qty": ln.get("qty"), "issued_qty": ln.get("issued_qty"),
                             "at": r.get("created_at"), "source": "Indent", "ref": r.get("code"), "status": r.get("status")})
    for m in (await active_medications(ctx, pid))[:10]:
        if m.get("encounter_id") in (None, aid):
            it = (await items_by_id(ctx, [m["item_id"]])).get(m["item_id"])
            meds.append({"name": item_label(it), "detail": " ".join(str(x) for x in (m.get("dose"), m.get("route"), m.get("frequency")) if x),
                         "at": m.get("prescribed_at"), "source": "Prescription", "ref": m.get("prescription_code"), "status": "Active"})
    meds.sort(key=lambda x: str(x.get("at") or ""), reverse=True)

    procs = (await _charge_rows(ctx, aid, scope="procedures", limit=8))
    nt = E(NOTES).table
    notes_ = await _rows(ctx, NOTES, nt.c.encounter_id == aid, order=_dcol(NOTES, "noted_at").desc().nulls_last(), limit=5)
    an = await _emp_names(ctx, [n.get("author") for n in notes_])
    return {
        "investigations": inv, "medications": meds[:8], "procedures": procs,
        "notes": [{"id": n["id"], "note_type": n.get("note_type"), "noted_at": n.get("noted_at"), "author": an.get(n.get("author") or ""),
                   "shift": n.get("shift"), "specialty": n.get("specialty"), "note": (n.get("note") or "")[:400], "status": n.get("status")}
                  for n in notes_],
        "tracker": await visit_tracker(ctx, adm),
    }


@router.get("/admissions/{aid}/timeline")
async def timeline(aid: str, ctx: CtxDep, limit: int = 150):
    rs.ensure_access(ctx, E(ADM), "read")
    adm = await _admission(ctx, aid)
    ev: list[dict] = []

    def add(at: Any, kind: str, title: str, detail: str | None = None, by: str | None = None, tone: str = "gray") -> None:
        if at:
            ev.append({"at": to_json_value(at) if not isinstance(at, str) else at, "kind": kind, "title": title, "detail": detail,
                       "by": by, "tone": tone})

    add(adm.get("started_at"), "admission", f"Admitted · {adm.get('code')}", adm.get("diagnosis"), adm.get("admitted_by"), "blue")
    add(adm.get("accepted_at"), "acceptance", "Accepted on the ward", adm.get("acceptance_remarks"), adm.get("accepted_by"), "green")
    for c in adm.get("condition_log") or []:
        add(c.get("at"), "condition", f"Condition: {c.get('previous')} → {c.get('condition')}", c.get("note"), c.get("by"),
            "red" if c.get("condition") == "Critical" else "amber" if c.get("condition") == "Serious" else "gray")
    for s in adm.get("sponsor_history") or []:
        add(s.get("at"), "sponsor", f"Sponsor changed to {(s.get('to') or {}).get('sponsor_name') or 'Self-pay'}", s.get("reason"), s.get("by"))
    for c in adm.get("patient_change_log") or []:
        add(c.get("at"), "patient", f"Patient re-linked {(c.get('from') or {}).get('uhid')} → {(c.get('to') or {}).get('uhid')}",
            c.get("reason"), c.get("by"), "red")
    tt = E(TRF).table
    for t in await _rows(ctx, TRF, tt.c.encounter_id == aid, limit=50):
        add(t.get("completed_at") or t.get("created_at"), "transfer", f"{t.get('transfer_type')} · {t.get('status')}", t.get("reason"))
    o = orders
    rows = (await ctx.conn.execute(_order_q().where(o.c.encounter_id == aid).order_by(o.c.ordered_at.desc()).limit(80))).all()
    users_, docs = await _order_people(ctx, rows)
    for r in rows:
        name = r.svc_name or r.svc_code
        add(r.ordered_at, "order", f"Ordered {name}", r.priority if r.priority and r.priority != "Routine" else None,
            docs.get((r.data or {}).get("ordering_doctor") or "") or users_.get(r.ordered_by or ""), "blue")
        if r.verified_at:
            add(r.verified_at, "result", f"Reported {name}", "Critical result" if r.critical else None, (r.data or {}).get("verified_by"),
                "red" if r.critical else "green")
        if r.status == "Cancelled":
            add((r.data or {}).get("voided_at") or (r.data or {}).get("cancelled_at"), "void", f"Cancelled {name}",
                (r.data or {}).get("cancel_reason"), (r.data or {}).get("cancelled_by"), "red")
    vt = E(VITALS).table
    vrows = await _rows(ctx, VITALS, vt.c.encounter_id == aid, sa.func.coalesce(vt.c.status, "") != "Entered in Error",
                        order=_dcol(VITALS, "recorded_at").desc(), limit=40)
    vn = await _emp_names(ctx, [v.get("recorded_by") for v in vrows])
    for v in vrows:
        n2 = v.get("news2")
        add(v.get("recorded_at"), "vitals", "Observations recorded", f"NEWS2 {n2} · {v.get('news2_risk')}" if n2 is not None else None,
            vn.get(v.get("recorded_by") or ""), "red" if (n2 or 0) >= 7 else "amber" if (n2 or 0) >= 5 else "gray")
    nt = E(NOTES).table
    nrows = await _rows(ctx, NOTES, nt.c.encounter_id == aid, order=_dcol(NOTES, "noted_at").desc(), limit=40)
    nn = await _emp_names(ctx, [n.get("author") for n in nrows])
    for n in nrows:
        add(n.get("noted_at"), "note", f"{n.get('note_type')}{' · ' + n['shift'] if n.get('shift') else ''}", (n.get("note") or "")[:140],
            nn.get(n.get("author") or ""))
    it = E(INSTR).table
    irows = await _rows(ctx, INSTR, it.c.encounter_id == aid, limit=40)
    gn = await _emp_names(ctx, [i.get("given_by") for i in irows])
    for i in irows:
        add(i.get("given_at"), "instruction", f"Doctor's instruction ({i.get('priority')})", (i.get("message") or "")[:140],
            gn.get(i.get("given_by") or ""), "amber" if i.get("priority") != "Routine" else "gray")
        add(i.get("acknowledged_at"), "instruction", "Instruction acknowledged", i.get("action_taken"), i.get("acknowledged_by"), "green")
    if has(INDENT):
        ti = E(INDENT).table
        for r in await _rows(ctx, INDENT, ti.c.encounter_id == aid, limit=40):
            add(r.get("created_at"), "indent", f"Medication indent {r.get('code')} · {r.get('status')}",
                f"{len(r.get('items') or [])} item(s)", r.get("raised_by"))
    if has(FORM_ENTRIES):
        ft = E(FORM_ENTRIES).table
        frows = await _rows(ctx, FORM_ENTRIES, ft.c.encounter_id == aid, ft.c.status.in_(["Signed", "Amended"]), limit=60)
        fn = await _emp_names(ctx, [f.get("author") for f in frows])
        for f in frows:
            band = f.get("risk_band")
            add(f.get("signed_at") or f.get("noted_at"), "form", f"{f.get('form_title') or f.get('form_key')} signed",
                " · ".join(str(x) for x in (f"{f.get('score_label') or 'Score'} {f.get('score')}" if f.get("score") is not None else None,
                                            band, f.get("summary")) if x) or None,
                fn.get(f.get("author") or ""), "red" if band and "high" in str(band).lower() else "gray")
    if has(CONSENT):
        ct = E(CONSENT).table
        for c in await _rows(ctx, CONSENT, ct.c.encounter_id == aid, limit=30):
            add(c.get("signed_at") or c.get("created_at"), "consent", f"Consent {'recorded' if c.get('status') == 'Signed' else 'not taken'}: "
                f"{c.get('consent_type')}", c.get("procedure_name"), c.get("signer_name"), "green" if c.get("status") == "Signed" else "amber")
    ut = E(USAGE).table
    for u in await _rows(ctx, USAGE, ut.c.encounter_id == aid, limit=40):
        add(u.get("connected_at"), "equipment", f"Connected {u.get('equipment_name')}", u.get("remarks"), u.get("connected_by"))
        add(u.get("disconnected_at"), "equipment", f"Disconnected {u.get('equipment_name')}",
            f"{u.get('units')} unit(s) · ₹{float(u.get('amount') or 0):,.2f}" if u.get("amount") is not None else None, u.get("disconnected_by"))
    add(adm.get("discharge_initiated_at"), "discharge", "Discharge initiated", adm.get("discharge_type"), adm.get("discharge_initiated_by"), "amber")
    ct2 = E(CLR).table
    for c in await _rows(ctx, CLR, ct2.c.encounter_id == aid, limit=30):
        add(c.get("cleared_at"), "clearance", f"{c.get('department_type')} cleared", c.get("remarks"), c.get("cleared_by"), "green")
    if adm.get("status") == "Discharged":
        add(adm.get("ended_at"), "discharge", "Discharged", adm.get("discharge_type"), None, "green")
    ev.sort(key=lambda x: str(x["at"]), reverse=True)
    return {"events": ev[: max(1, min(limit, 500))]}


# ================================================================== progress notes, templates, doctor's instructions

@router.get("/note-templates")
async def note_templates(ctx: CtxDep, kind: Literal["clinical", "nursing"] = "clinical"):
    return {"templates": NOTE_TEMPLATES[kind]}


@router.get("/admissions/{aid}/notes")
async def notes(aid: str, ctx: CtxDep, kind: Literal["clinical", "nursing", "all"] = "all", shift: str | None = None, limit: int = 100):
    rs.ensure_access(ctx, E(NOTES), "read")
    await _admission(ctx, aid)
    nt = E(NOTES).table
    rows = await _rows(ctx, NOTES, nt.c.encounter_id == aid, order=_dcol(NOTES, "noted_at").desc().nulls_last(), limit=max(1, min(limit, 500)))
    if kind == "clinical":
        rows = [r for r in rows if r.get("note_type") != "Nursing Note"]
    elif kind == "nursing":
        rows = [r for r in rows if r.get("note_type") == "Nursing Note"]
    if shift:
        rows = [r for r in rows if r.get("shift") == shift]
    names = await _emp_names(ctx, [r.get("author") for r in rows] + [r.get("cosigner") for r in rows])
    return {"rows": [{"id": r["id"], "code": r.get("code"), "note_type": r.get("note_type"), "noted_at": r.get("noted_at"),
                      "author": names.get(r.get("author") or ""), "author_role": r.get("author_role"), "shift": r.get("shift"),
                      "specialty": r.get("specialty"), "note": r.get("note"), "status": r.get("status"), "signed_at": r.get("signed_at"),
                      "amended_at": r.get("amended_at")} for r in rows]}


class NoteIn(BaseModel):
    kind: Literal["clinical", "nursing"] = "clinical"
    note_type: Literal["Progress Note", "Ward Round", "Consult Note", "Procedure Note", "Nursing Note"] | None = None
    shift: Literal["Morning", "Evening", "Night"]
    specialty: str | None = Field(default=None, max_length=80)
    note: str = Field(min_length=3, max_length=20000)


def _current_shift() -> str:
    h = now().hour
    return "Morning" if 8 <= h < 14 else "Evening" if 14 <= h < 20 else "Night"


@router.post("/admissions/{aid}/notes", status_code=201)
async def add_note(aid: str, body: NoteIn, ctx: CtxDep):
    if body.kind == "clinical":
        _need(ctx, "doctor_note", "write a doctor's progress note")
        note_type = body.note_type if body.note_type and body.note_type != "Nursing Note" else "Progress Note"
    else:
        _need(ctx, "nursing_note", "write a nursing note")
        note_type = "Nursing Note"
    adm = await _admission(ctx, aid)
    _require_open(adm, "add notes")
    rec = await create(ctx, NOTES, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "note_type": note_type, "noted_at": _stamp(), "author": _emp(ctx),
        "author_role": ctx.role, "department": adm.get("department_id"), "note": body.note.strip(), "shift": body.shift,
        "specialty": (body.specialty or "").strip() or None, "status": "Signed", "signed_at": _stamp(), "source": "ipchart",
    })
    return {"id": rec["id"], "code": rec.get("code"), "note_type": note_type, "noted_at": rec.get("noted_at"), "shift": body.shift}


FOR_ROLE_ROLES = {"Nursing": {"nurse"}, "Dietitian": {"nurse", "support"}, "Physiotherapy": {"nurse", "support"},
                  "Pharmacy": {"pharmacist"}, "Billing / front office": {"billing", "front_office"}, "Resident doctor": {"doctor"}}


@router.get("/admissions/{aid}/instructions")
async def instructions(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(INSTR), "read")
    await _admission(ctx, aid)
    it = E(INSTR).table
    rows = await _rows(ctx, INSTR, it.c.encounter_id == aid, order=[it.c.status.desc(), _dcol(INSTR, "given_at").desc()], limit=300)
    names = await _emp_names(ctx, [r.get("given_by") for r in rows])
    role = ctx.role
    out = [{"id": r["id"], "code": r.get("code"), "message": r.get("message"), "priority": r.get("priority"), "for_role": r.get("for_role"),
            "given_by": names.get(r.get("given_by") or ""), "given_at": r.get("given_at"), "status": r.get("status"),
            "acknowledged_at": r.get("acknowledged_at"), "acknowledged_by": r.get("acknowledged_by"), "action_taken": r.get("action_taken"),
            "can_ack": r.get("status") == "Open" and (role == "admin" or role in FOR_ROLE_ROLES.get(r.get("for_role") or "Nursing", {"nurse"}))}
           for r in rows]
    out.sort(key=lambda x: (x["status"] != "Open", {"STAT": 0, "Urgent": 1}.get(x["priority"] or "", 2), str(x["given_at"] or "")), reverse=False)
    return {"rows": out, "open": sum(1 for x in out if x["status"] == "Open")}


class InstructionIn(BaseModel):
    message: str = Field(min_length=3, max_length=2000)
    priority: Literal["Routine", "Urgent", "STAT"] = "Routine"
    for_role: Literal["Nursing", "Dietitian", "Physiotherapy", "Pharmacy", "Billing / front office", "Resident doctor"] = "Nursing"


@router.post("/admissions/{aid}/instructions", status_code=201)
async def add_instruction(aid: str, body: InstructionIn, ctx: CtxDep):
    _need(ctx, "instruct", "leave doctor's instructions")
    adm = await _admission(ctx, aid)
    _require_open(adm, "leave instructions")
    rec = await create(ctx, INSTR, {"patient_id": adm["patient_id"], "encounter_id": aid, "message": body.message.strip(),
                                    "priority": body.priority, "for_role": body.for_role, "given_by": _emp(ctx), "given_at": _stamp(),
                                    "status": "Open"})
    bed = await _bed_info(ctx, adm.get("bed_id"))
    live(ctx, "alerts", {"type": "doctor_instruction", "roles": sorted(FOR_ROLE_ROLES[body.for_role]), "admission_id": aid,
                         "message": f"{'STAT ' if body.priority == 'STAT' else ''}Instruction from {_who(ctx)} for {adm.get('code')}"
                                    f"{' · bed ' + str(bed['bed_no']) if bed else ''}: {body.message[:90]}"})
    return rec


class AckIn(BaseModel):
    action_taken: str | None = Field(default=None, max_length=500)


@router.post("/instructions/{iid}/acknowledge")
async def acknowledge_instruction(iid: str, body: AckIn, ctx: CtxDep):
    ins = await get(ctx, INSTR, iid)
    if not ins:
        raise not_found("Instruction not found")
    allowed = FOR_ROLE_ROLES.get(ins.get("for_role") or "Nursing", {"nurse"})
    if not _can(ctx, *allowed):
        raise forbidden(f"Instructions for {ins.get('for_role')} are acknowledged by {', '.join(sorted(allowed))}")
    if ins.get("status") != "Open":
        raise conflict(f"Already {str(ins.get('status')).lower()}")
    values = {"action_taken": (body.action_taken or "").strip() or None}
    return await transition(ctx, INSTR, iid, "Acknowledged", values={k: v for k, v in values.items() if v} or None)


@hook(INSTR, "before_transition")
async def _instruction_ack(ctx, nxt: dict, prev: dict) -> None:
    if nxt["status"] == "Acknowledged":
        nxt["acknowledged_by"] = f"{_who(ctx)} ({ctx.role})"


@router.post("/instructions/{iid}/withdraw")
async def withdraw_instruction(iid: str, ctx: CtxDep):
    _need(ctx, "instruct", "withdraw instructions")
    ins = await get(ctx, INSTR, iid)
    if not ins:
        raise not_found("Instruction not found")
    if ins.get("status") != "Open":
        raise conflict(f"Already {str(ins.get('status')).lower()}")
    return await transition(ctx, INSTR, iid, "Cancelled", note=f"Withdrawn by {_who(ctx)}")


# ================================================================== consents

@router.get("/admissions/{aid}/consents")
async def consents(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(CONSENT), "read")
    await _admission(ctx, aid)
    ct = E(CONSENT).table
    rows = await _rows(ctx, CONSENT, ct.c.encounter_id == aid, order=ct.c.created_at.desc(), limit=200)
    names = await _emp_names(ctx, [r.get("explained_by") for r in rows])
    tpls = await find(ctx, CONSENT_TPL, status="Active", limit=200, sort="title")
    fm = E(CONSENT).field_map
    return {
        "rows": [{"id": r["id"], "code": r.get("code"), "status": r.get("status"), "taken": r.get("status") == "Signed",
                  "consent_type": r.get("consent_type"), "template": r.get("template_name"), "procedure": r.get("procedure_name"),
                  "signer_name": r.get("signer_name"), "relationship": r.get("signer_relationship"), "witness": r.get("witness_name"),
                  "explained_by": names.get(r.get("explained_by") or "") or r.get("explained_by_name"), "language": r.get("language"),
                  "risks": r.get("risks_explained"), "benefits": r.get("benefits_explained"), "alternatives": r.get("alternatives_explained"),
                  "remarks": r.get("remarks"), "signed_at": r.get("signed_at"), "recorded_by": r.get("recorded_by"),
                  "recorded_at": r.get("created_at"), "has_signature": bool(r.get("signature"))} for r in rows],
        "templates": [{"id": t["id"], "name": t.get("name"), "consent_type": t.get("consent_type"), "version": t.get("version")} for t in tpls],
        "types": fm["consent_type"].options, "relationships": fm["signer_relationship"].options, "languages": fm["language"].options,
    }


class ConsentIn(BaseModel):
    taken: bool = True
    template_id: str | None = None
    consent_type: str | None = None
    procedure: str | None = Field(default=None, max_length=200)
    signer_name: str = Field(min_length=2, max_length=120)
    relationship: str = "Self"
    witness: str | None = Field(default=None, max_length=120)
    explained_by: str | None = None
    language: str | None = None
    risks: str | None = Field(default=None, max_length=2000)
    benefits: str | None = Field(default=None, max_length=2000)
    alternatives: str | None = Field(default=None, max_length=2000)
    remarks: str | None = Field(default=None, max_length=2000)
    signature: str | None = None


async def _general_template(ctx) -> dict:
    """The active general-admission consent template, created (v1.0, built-in text) on first use."""
    for t in await find(ctx, CONSENT_TPL, status="Active", consent_type="General", limit=20, sort="title"):
        return t
    return await create(ctx, CONSENT_TPL, {
        "name": "General admission consent", "consent_type": "General", "version": "1.0", "body": GENERAL_CONSENT_TEXT,
        "languages": ["English"], "requires_witness": True, "requires_doctor_signature": True, "effective_from": now().date().isoformat(),
        "published_at": _stamp(), "change_note": "Built-in default — review and replace with the hospital's approved text", "builtin": True,
        "status": "Active",
    })


@router.post("/admissions/{aid}/consents", status_code=201)
async def record_consent(aid: str, body: ConsentIn, ctx: CtxDep):
    """Structured consent recording (taken / not taken) stored in Consent Management against this admission."""
    _need(ctx, "consent", "record consents")
    adm = await _admission(ctx, aid)
    fm = E(CONSENT).field_map
    if body.relationship not in (fm["signer_relationship"].options or []):
        raise bad_request("Choose the signatory's relationship", {"fields": {"relationship": "Invalid"}})
    if body.language and body.language not in (fm["language"].options or []):
        raise bad_request("Choose the language from the list", {"fields": {"language": "Invalid"}})
    if body.consent_type and body.consent_type not in (fm["consent_type"].options or []):
        raise bad_request("Unknown consent type", {"fields": {"consent_type": "Invalid"}})
    if not body.taken and not (body.remarks or "").strip():
        raise bad_request("Record why consent was not taken", {"fields": {"remarks": "Reason required"}})
    if body.signature and (not body.signature.startswith("data:image/") or len(body.signature) > 800_000):
        raise bad_request("Invalid signature image")
    tpl = None
    if body.template_id:
        tpl = await get(ctx, CONSENT_TPL, body.template_id)
        if not tpl or tpl.get("status") != "Active":
            raise bad_request("Choose an active consent template", {"fields": {"template_id": "Inactive / unknown"}})
    elif body.consent_type:
        tpl = await first(ctx, CONSENT_TPL, status="Active", consent_type=body.consent_type)
    tpl = tpl or await _general_template(ctx)
    explained = body.explained_by or (_emp(ctx) if ctx.role == "doctor" else None)
    if explained:
        await _doctor(ctx, explained, "explained_by")
    rec = await create(ctx, CONSENT, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "template": tpl["id"], "template_name": tpl.get("name"),
        "consent_type": body.consent_type or tpl.get("consent_type"), "procedure_name": (body.procedure or "").strip() or None,
        "explained_by": explained, "risks_explained": body.risks, "benefits_explained": body.benefits,
        "alternatives_explained": body.alternatives, "remarks": body.remarks, "language": body.language or "English",
        "signer_name": body.signer_name.strip(), "signer_relationship": body.relationship, "witness_name": (body.witness or "").strip() or None,
        "signature": body.signature, "consent_mode": "Tap-to-sign (tablet)" if body.signature else "Paper (scanned)",
        "consent_taken": body.taken, "recorded_by": _who(ctx), "status": "Signed" if body.taken else "Pending",
        "signed_at": _stamp() if body.taken else None,
    })
    return {"id": rec["id"], "code": rec.get("code"), "status": rec.get("status"), "template": tpl.get("name")}


# ================================================================== order desk — catalogue & money helpers

@router.get("/services")
async def service_catalogue(ctx: CtxDep, kind: Literal["investigation", "service", "procedure", "grade"] = "investigation",
                            q: str | None = None, limit: int = 40):
    rs.ensure_access(ctx, E(ADM), "read")
    cats = {"investigation": INVESTIGATION_CATS, "service": tuple(SERVICE_REQUEST_CATS), "procedure": CHARGE_CATS, "grade": GRADE_CATS}[kind]
    s = services
    conds = [s.c.deleted_at.is_(None), s.c.category.in_(cats), sa.func.coalesce(s.c.status, "Active") == "Active"]
    if kind in ("service", "procedure"):  # per-day / home-care heads are posted by their own workflows
        conds += [sa.not_(s.c.code.ilike("%DAILY%")), sa.not_(s.c.code.ilike("HOME-%")), sa.not_(s.c.code.ilike("MORTUARY%"))]
    if q and q.strip():
        term = f"%{q.strip().lower()}%"
        conds.append(sa.or_(sa.func.lower(s.c.title).like(term), sa.func.lower(s.c.code).like(term)))
    rows = (await ctx.conn.execute(sa.select(s).where(*conds).order_by(s.c.title).limit(max(1, min(limit, 100))))).all()
    money_ok = _money_ok(ctx)
    out = []
    for r in rows:
        rec = rs.to_record(r)
        x = {"id": rec["id"], "code": rec.get("code"), "name": rec.get("name") or rec.get("title"), "category": rec.get("category"),
             "tat_minutes": rec.get("tat_minutes"), "sample_type": rec.get("sample_type"), "modality": rec.get("modality")}
        if money_ok:
            x.update({k: rec.get(k) for k in ("rate", "rate_private", "rate_icu", "rate_international", "tax_percent")})
        out.append(x)
    return {"rows": out}


def _amounts(rate: Decimal, qty: Decimal, discount: float, mode: str) -> tuple[Decimal, Decimal, Decimal]:
    gross = money(rate * qty)
    disc = money(gross * Decimal(str(discount)) / 100) if mode == "percent" else money(discount)
    if disc > gross:
        raise bad_request("Discount cannot exceed the line amount")
    return gross, disc, money(gross - disc)


def _check_discount(ctx, gross: Decimal, disc: Decimal, reason: str | None) -> None:
    if disc <= 0:
        return
    if not _perm(ctx, "discount"):
        raise forbidden("Only billing / front office can give a discount")
    if len((reason or "").strip()) < 3:
        raise bad_request("Give a reason for the discount", {"fields": {"discount_reason": "Required"}})
    pct = float(disc / gross * 100) if gross else 0
    if pct > settings.discount_approval_percent and not _can(ctx, "billing"):
        raise forbidden(f"Discounts above {settings.discount_approval_percent:g}% need a billing user")


async def _ordering_doctor(ctx, adm: dict, given: str | None) -> tuple[str, bool]:
    if ctx.role == "doctor" and _emp(ctx):
        return _emp(ctx), False
    did = given or adm.get("doctor_id")
    await _doctor(ctx, did, "ordering_doctor")
    return did, True


async def _charge_of(ctx, source_id: str) -> dict | None:
    row = (await ctx.conn.execute(sa.select(charges.c.id).where(charges.c.source_id == source_id, charges.c.deleted_at.is_(None))
                                  .order_by(charges.c.created_at.desc()).limit(1))).first()
    return await get(ctx, CHARGES, str(row.id)) if row else None


async def _unbilled_total(ctx, aid: str) -> Decimal:
    v = (await ctx.conn.execute(sa.select(sa.func.coalesce(sa.func.sum(charges.c.amount), 0)).where(
        charges.c.encounter_id == aid, charges.c.status == "Unbilled", charges.c.deleted_at.is_(None)))).scalar_one()
    return Decimal(v)


async def _advance(ctx, adm: dict) -> Decimal:
    """Advance left after the running (unbilled) bill — negative means the patient owes."""
    return await deposit_balance(ctx, adm["patient_id"]) - await _unbilled_total(ctx, adm["id"])


# ================================================================== investigations

class InvLine(BaseModel):
    service_id: str
    qty: int = Field(default=1, ge=1, le=20)
    discount: float = Field(default=0, ge=0)
    discount_mode: Literal["amount", "percent"] = "amount"
    priority: Literal["Routine", "Urgent", "STAT"] = "Routine"


class InvOrderIn(BaseModel):
    items: list[InvLine] = Field(min_length=1, max_length=30)
    clinical_notes: str | None = Field(default=None, max_length=500)
    ordering_doctor: str | None = None
    discount_reason: str | None = Field(default=None, max_length=300)
    confirm_duplicates: bool = False


async def _recent_duplicates(ctx, pid: str, service_ids: list[str]) -> list[dict]:
    o, s = orders, services
    rows = (await ctx.conn.execute(
        sa.select(o.c.id, o.c.code, o.c.status, o.c.ordered_at, o.c.service_id, s.c.title)
        .select_from(o.outerjoin(s, s.c.id == o.c.service_id))
        .where(o.c.patient_id == pid, o.c.service_id.in_(service_ids), o.c.ordered_at >= now() - timedelta(hours=24),
               o.c.status != "Cancelled", o.c.deleted_at.is_(None)).order_by(o.c.ordered_at.desc()))).all()
    return [{"service_id": str(r.service_id), "test": r.title, "order_id": str(r.id), "code": r.code, "status": r.status,
             "ordered_at": _iso(r.ordered_at)} for r in rows]


async def _services_by_id(ctx, ids: list[str]) -> dict[str, dict]:
    ids = [i for i in ids if is_uuid(i)]
    if not ids:
        return {}
    rows = (await ctx.conn.execute(sa.select(services).where(services.c.id.in_(ids), services.c.deleted_at.is_(None)))).all()
    return {str(r.id): rs.to_record(r) for r in rows}


@router.post("/admissions/{aid}/investigations", status_code=201)
async def order_investigations(aid: str, body: InvOrderIn, ctx: CtxDep):
    """Order tests / scans for the stay. Lines share one indent number; each order's charge carries the
    quantity and any line discount (charge.amount is net of discount; gross kept in data.gross_amount)."""
    _need(ctx, "order", "order investigations")
    adm = await _admission(ctx, aid)
    _require_open(adm, "order investigations")
    ids = [i.service_id for i in body.items]
    if len(set(ids)) != len(ids):
        raise bad_request("The same test is listed twice")
    svcs = await _services_by_id(ctx, ids)
    bad = [i for i in ids if i not in svcs or svcs[i].get("category") not in INVESTIGATION_CATS or svcs[i].get("status") == "Inactive"]
    if bad:
        raise bad_request("Not orderable as an investigation: " + ", ".join(svcs[i].get("name", "?") if i in svcs else "unknown" for i in bad))
    dups = await _recent_duplicates(ctx, adm["patient_id"], ids)
    if dups and not body.confirm_duplicates:
        raise ApiError(409, f"{len({d['service_id'] for d in dups})} test(s) already ordered in the last 24 hours", {"warnings": dups},
                       code="DUPLICATE_ORDER")
    doctor, verbal = await _ordering_doctor(ctx, adm, body.ordering_doctor)
    lines = []
    for it in body.items:
        rate = Decimal(str(svcs[it.service_id].get("rate") or 0))
        gross, disc, net = _amounts(rate, Decimal(it.qty), it.discount, it.discount_mode)
        _check_discount(ctx, gross, disc, body.discount_reason)
        lines.append((it, rate, gross, disc, net))
    indent_no = await next_code(ctx.conn, "INDT")
    batch = new_id()
    warned = {d["service_id"] for d in dups}
    p = await _patient_row(ctx, adm["patient_id"])
    created = []
    for it, rate, gross, disc, net in lines:
        svc = svcs[it.service_id]
        rec = await create(ctx, ORD, {
            "patient_id": adm["patient_id"], "encounter_id": aid, "service_id": it.service_id, "category": INVESTIGATION_TO_ORDER[svc["category"]],
            "priority": it.priority, "clinical_notes": (body.clinical_notes or "").strip() or None, "ordering_doctor": doctor,
            "status": "Ordered", "order_batch": batch, "indent_no": indent_no, "qty": it.qty, "verbal_order": verbal,
            "entered_by": _who(ctx), "duplicate_override": it.service_id in warned, "source": "ipchart.order_desk",
        })
        ch = await _charge_of(ctx, rec["id"])
        if ch:
            await update(ctx, CHARGES, ch["id"], {"qty": it.qty, "rate": float(rate), "amount": float(net), "discount": float(disc),
                                                  "gross_amount": float(gross), "indent_no": indent_no,
                                                  "discount_reason": (body.discount_reason or "").strip() or None if disc else None})
        created.append({"id": rec["id"], "code": rec["code"], "test": svc.get("name"), "category": rec.get("category"), "qty": it.qty,
                        "priority": it.priority})
        if it.priority == "STAT":
            live(ctx, "alerts", {"type": "stat_order", "roles": ["lab", "radiology", "nurse"], "category": rec.get("category"),
                                 "order_id": rec["id"], "message": f"STAT {svc.get('name')} for {p.get('title')} ({adm.get('code')})"})
    gross_t = sum((x[2] for x in lines), Decimal(0))
    disc_t = sum((x[3] for x in lines), Decimal(0))
    if disc_t:
        await audit(ctx, "ipchart.discount", module="orders", entity="orders", record={"id": created[0]["id"], "patient_id": adm["patient_id"]},
                    after={"indent_no": indent_no, "gross": float(gross_t), "discount": float(disc_t)}, note=body.discount_reason)
    out = {"indent_no": indent_no, "orders": created, "count": len(created), "warnings": dups}
    if _money_ok(ctx):
        out.update(gross=float(gross_t), discount=float(disc_t), total=float(gross_t - disc_t))
    return out


async def _indent_payments(ctx, adm: dict, order_rows: list) -> list[dict]:
    by_indent: dict[str, dict] = {}
    for r in order_rows:
        d = r["data"] or {}
        code = d.get("indent_no")
        if not code:
            continue
        x = by_indent.setdefault(code, {"indent_no": code, "billed": Decimal(0), "orders": 0, "cancelled": 0, "at": r["ordered_at"]})
        x["orders"] += 1
        if r["status"] == "Cancelled":
            x["cancelled"] += 1
        else:
            x["billed"] += Decimal(str(r["amount"] or 0))
    if not by_indent:
        return []
    pay = payments
    key = pay.c.data.op("->>")(sa.literal_column("'indent_no'")).label("indent_no")
    paid = {k: Decimal(v) for k, v in (await ctx.conn.execute(
        sa.select(key, sa.func.coalesce(sa.func.sum(pay.c.amount), 0))
        .where(pay.c.patient_id == adm["patient_id"], pay.c.deleted_at.is_(None), pay.c.status == "Success",
               pay.c.kind.in_(["DEPOSIT", "PAYMENT"]), key.element.in_(list(by_indent)))
        .group_by(key.element))).all()}
    adv = await _advance(ctx, adm)
    sponsored = (adm.get("payer") or "Cash") != "Cash"
    out = []
    for code, x in sorted(by_indent.items(), key=lambda kv: str(kv[1]["at"]), reverse=True):
        coll = paid.get(code, Decimal(0))
        if x["orders"] == x["cancelled"]:
            st = "Cancelled"
        elif coll >= x["billed"] > 0 or (x["billed"] == 0 and coll > 0):
            st = "Paid"
        elif sponsored:
            st = "Sponsor credit"
        elif adv >= 0:
            st = "Covered by advance"
        else:
            st = "Payment due"
        out.append({"indent_no": code, "billed": float(x["billed"]), "collected": float(coll), "due": float(max(x["billed"] - coll, Decimal(0))),
                    "status": st, "orders": x["orders"], "at": _iso(x["at"])})
    return out


@router.get("/admissions/{aid}/investigations")
async def investigations(aid: str, ctx: CtxDep, scope: Literal["investigations", "services", "all"] = "investigations"):
    if not (_ent_ok(ctx, ORD) or _perm(ctx, "collect")):
        raise forbidden("Your role cannot view investigations")
    adm = await _admission(ctx, aid)
    o, c = orders, charges
    q = (_order_q().add_columns(c.c.qty.label("ch_qty"), c.c.rate.label("ch_rate"), c.c.amount.label("ch_amount"),
                                c.c.discount.label("ch_discount"), c.c.status.label("ch_status"))
         .outerjoin(c, sa.and_(c.c.source_id == o.c.id, c.c.deleted_at.is_(None))).where(o.c.encounter_id == aid))
    if scope == "investigations":
        q = q.where(o.c.category.in_(INVESTIGATION_CATS), sa.func.coalesce(o.c.data["kind"].astext, "") != "service_request")
    elif scope == "services":
        q = q.where(sa.or_(o.c.data["kind"].astext == "service_request", o.c.category.in_(["NURSING", "DIET"])))
    rows = (await ctx.conn.execute(q.order_by(o.c.ordered_at.desc().nulls_last()).limit(1000))).all()
    users_, docs = await _order_people(ctx, rows)
    money_ok = _money_ok(ctx)
    out, raw = [], []
    for r in rows:
        d = r.data or {}
        amt = r.ch_amount if r.ch_status != "Cancelled" else Decimal(0)
        raw.append({"data": d, "status": r.status, "amount": amt or 0, "ordered_at": r.ordered_at})
        x = {"id": str(r.id), "code": r.code, "test": r.svc_name or r.svc_code, "category": r.category, "status": r.status,
             "stage": _stage(r.status), "priority": r.priority or "Routine", "indent_no": d.get("indent_no"),
             "qty": float(r.ch_qty) if r.ch_qty is not None else d.get("qty") or 1, "ordered_at": _iso(r.ordered_at),
             "ordering_doctor": docs.get(d.get("ordering_doctor") or ""), "created_by": users_.get(r.ordered_by or "") or d.get("entered_by"),
             "critical": bool(r.critical), "report_ready": r.status in ("Resulted", "Verified"), "verified_at": _iso(r.verified_at),
             "voided_by": d.get("voided_by") or (d.get("cancelled_by") if r.status == "Cancelled" else None),
             "voided_at": d.get("voided_at"), "void_reason": d.get("cancel_reason"), "charge_status": r.ch_status,
             "can_void": r.status in ("Ordered", "Collected") and _perm(ctx, "void_order"), "notes": d.get("clinical_notes")}
        if money_ok:
            x.update(rate=float(r.ch_rate) if r.ch_rate is not None else float(r.svc_rate or 0),
                     discount=float(r.ch_discount or 0), total=float(r.ch_amount or 0))
        out.append(x)
    res: dict[str, Any] = {"rows": out, "counts": {s: sum(1 for x in out if x["stage"] == s) for s in ("Ordered", "Sample collected", "In progress", "Reported", "Cancelled")}}
    if money_ok:
        res["indents"] = await _indent_payments(ctx, adm, raw)
        res["advance_available"] = float(await _advance(ctx, adm))
        res["payer"] = adm.get("payer") or "Cash"
    return res


class CollectIn(BaseModel):
    indent_no: str = Field(min_length=3, max_length=40)
    mode: Literal["CASH", "CARD", "UPI", "WALLET", "BANK", "CHEQUE"]
    amount: float = Field(gt=0, le=10_000_000)
    reference: str | None = Field(default=None, max_length=80)


@router.post("/admissions/{aid}/collect", status_code=201)
async def collect_payment(aid: str, body: CollectIn, ctx: CtxDep):
    """Collect payment for an investigations indent (cash patients) — booked as an advance against the stay."""
    _need(ctx, "collect", "collect payments")
    adm = await _admission(ctx, aid)
    _require_open(adm, "collect payments")
    o = orders
    rows = (await ctx.conn.execute(sa.select(o.c.id, o.c.status, o.c.data).where(
        o.c.encounter_id == aid, o.c.deleted_at.is_(None), o.c.data["indent_no"].astext == body.indent_no))).all()
    if not rows:
        raise not_found(f"Indent {body.indent_no} is not on this admission")
    raw = []
    for r in rows:
        ch = await _charge_of(ctx, str(r.id))
        raw.append({"data": r.data, "status": r.status, "amount": (ch or {}).get("amount") or 0 if (ch or {}).get("status") != "Cancelled" else 0,
                    "ordered_at": None})
    info = next(iter(await _indent_payments(ctx, adm, raw)), None)
    if not info or info["status"] == "Cancelled":
        raise conflict(f"Indent {body.indent_no} is cancelled")
    if Decimal(str(body.amount)) > Decimal(str(info["due"])) + Decimal("0.01"):
        raise bad_request(f"Amount exceeds the ₹{info['due']:,.2f} due on {body.indent_no}")
    if body.mode != "CASH" and not (body.reference or "").strip():
        raise bad_request("Enter the transaction / cheque reference", {"fields": {"reference": "Required"}})
    pay = await create(ctx, PAYMENTS, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "kind": "DEPOSIT", "mode": body.mode, "amount": float(money(body.amount)),
        "reference": (body.reference or "").strip() or f"Investigations {body.indent_no}", "paid_at": _stamp(), "received_by": _who(ctx),
        "indent_no": body.indent_no, "purpose": "Investigations advance", "status": "Success",
    })
    await emit("payment.received", ctx, {"payment": pay})
    return {"payment": {"id": pay["id"], "code": pay.get("code"), "amount": pay.get("amount"), "mode": pay.get("mode")},
            "indent": {**info, "collected": info["collected"] + body.amount, "due": max(0.0, info["due"] - body.amount)}}


# ================================================================== service requests

class ServiceReqIn(BaseModel):
    service_id: str
    qty: int = Field(default=1, ge=1, le=60)
    priority: Literal["Routine", "Urgent", "STAT"] = "Routine"
    notes: str | None = Field(default=None, max_length=500)
    scheduled_at: str | None = None
    ordering_doctor: str | None = None


@router.post("/admissions/{aid}/service-requests", status_code=201)
async def service_request(aid: str, body: ServiceReqIn, ctx: CtxDep):
    """Physiotherapy, dietitian, nursing procedures, dialysis… raised as orders for the department to act on."""
    _need(ctx, "order", "raise service requests")
    adm = await _admission(ctx, aid)
    _require_open(adm, "raise service requests")
    svc = (await _services_by_id(ctx, [body.service_id])).get(body.service_id)
    if not svc or svc.get("category") not in SERVICE_REQUEST_CATS or svc.get("status") == "Inactive":
        raise bad_request("Choose a service from the list", {"fields": {"service_id": "Not a requestable service"}})
    sched = parse_dt(body.scheduled_at) if body.scheduled_at else None
    doctor, verbal = await _ordering_doctor(ctx, adm, body.ordering_doctor)
    rec = await create(ctx, ORD, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "service_id": body.service_id, "category": SERVICE_REQUEST_CATS[svc["category"]],
        "priority": body.priority, "clinical_notes": (body.notes or "").strip() or None, "ordering_doctor": doctor, "status": "Ordered",
        "kind": "service_request", "qty": body.qty, "scheduled_at": sched.isoformat() if sched else None, "verbal_order": verbal,
        "entered_by": _who(ctx), "source": "ipchart.order_desk",
    })
    ch = await _charge_of(ctx, rec["id"])
    if ch and body.qty != 1:
        rate = Decimal(str(ch.get("rate") or 0))
        await update(ctx, CHARGES, ch["id"], {"qty": body.qty, "amount": float(money(rate * body.qty)), "gross_amount": float(money(rate * body.qty))})
    return {"id": rec["id"], "code": rec.get("code"), "service": svc.get("name"), "status": rec.get("status")}


# ================================================================== procedure / grade / manual equipment charges

class ChargeIn(BaseModel):
    kind: Literal["procedure", "grade", "equipment"]
    service_id: str | None = None
    description: str | None = Field(default=None, max_length=200)
    qty: float = Field(default=1, gt=0, le=1000)
    rate: float | None = Field(default=None, ge=0, le=10_000_000)
    discount: float = Field(default=0, ge=0)
    discount_mode: Literal["amount", "percent"] = "amount"
    discount_reason: str | None = Field(default=None, max_length=300)
    grade: Literal["General", "Private", "ICU", "International"] | None = None
    doctor_id: str | None = None
    performed_at: str | None = None
    remarks: str | None = Field(default=None, max_length=300)


def _grade_for(bed: dict | None) -> str:
    cat = (bed or {}).get("category")
    return "ICU" if cat in ICU_BEDS else "Private" if cat in PRIVATE_BEDS else "General"


SOURCE_TYPE = {"ipchart.procedure": "Procedure", "ipchart.grade": "Visit / grade", "ipchart.equipment": "Equipment",
               USAGE: "Equipment", ORD: "Service order", CONSULTS: "Consultation", "billing.manual": "Added on bill"}


@router.post("/admissions/{aid}/charges", status_code=201)
async def add_charge(aid: str, body: ChargeIn, ctx: CtxDep):
    _need(ctx, "charge", "add charges")
    adm = await _admission(ctx, aid)
    _require_open(adm, "add charges")
    svc = (await _services_by_id(ctx, [body.service_id])).get(body.service_id or "") if body.service_id else None
    if body.service_id and not svc:
        raise bad_request("Unknown service", {"fields": {"service_id": "Not in the tariff"}})
    if body.rate is not None and svc and not _perm(ctx, "rate_override"):
        raise forbidden("Only billing can override the tariff rate")
    grade = None
    if body.kind == "procedure":
        if not svc or svc.get("category") not in CHARGE_CATS:
            raise bad_request("Choose a procedure / service from the tariff", {"fields": {"service_id": "Required"}})
        rate = Decimal(str(body.rate if body.rate is not None else svc.get("rate") or 0))
    elif body.kind == "grade":
        if not svc or svc.get("category") not in GRADE_CATS:
            raise bad_request("Choose the service to bill by grade", {"fields": {"service_id": "Required"}})
        grade = body.grade or _grade_for(await _bed_info(ctx, adm.get("bed_id")))
        graded = svc.get(GRADES[grade])
        rate = Decimal(str(body.rate if body.rate is not None else graded if graded not in (None, "") else svc.get("rate") or 0))
    else:
        if not svc and not (body.description or "").strip():
            raise bad_request("Name the equipment / item", {"fields": {"description": "Required"}})
        if body.rate is None and not svc:
            raise bad_request("Enter the rate", {"fields": {"rate": "Required"}})
        rate = Decimal(str(body.rate if body.rate is not None else svc.get("rate") or 0))
    doctor = await _doctor(ctx, body.doctor_id, "doctor_id") if body.doctor_id else None
    gross, disc, net = _amounts(rate, Decimal(str(body.qty)), body.discount, body.discount_mode)
    _check_discount(ctx, gross, disc, body.discount_reason)
    desc = (body.description or "").strip() or (svc or {}).get("name") or "Charge"
    if body.kind == "grade":
        desc = f"{desc} — {grade} grade" + (f" · {doctor['name']}" if doctor else "")
    elif body.kind == "equipment":
        desc = f"Equipment: {desc}"
    performed = parse_dt(body.performed_at) if body.performed_at else None
    ch = await post_charge(ctx, patient_id=adm["patient_id"], encounter_id=aid, service_id=(svc or {}).get("id"), qty=body.qty,
                           rate=float(rate), description=desc, source=f"ipchart.{body.kind}")
    ch = await update(ctx, CHARGES, ch["id"], {"amount": float(net), "discount": float(disc), "gross_amount": float(gross), "grade": grade,
                                               "performed_by": doctor["id"] if doctor else None, "performed_by_name": doctor["name"] if doctor else None,
                                               "performed_at": performed.isoformat() if performed else None, "remarks": body.remarks,
                                               "entered_by": _who(ctx), "discount_reason": (body.discount_reason or "").strip() or None if disc else None})
    return {"id": ch["id"], "code": ch.get("code"), "description": ch.get("description"), "qty": ch.get("qty"),
            **({"rate": ch.get("rate"), "amount": ch.get("amount"), "discount": ch.get("discount")} if _money_ok(ctx) else {})}


async def _charge_rows(ctx, aid: str, *, scope: str = "procedures", limit: int = 300) -> list[dict]:
    c, s = charges, services
    q = (sa.select(c.c.id, c.c.code, c.c.status, c.c.qty, c.c.rate, c.c.amount, c.c.discount, c.c.source, c.c.source_id, c.c.charged_at,
                   c.c.created_by, c.c.data, s.c.title.label("svc_name"), s.c.category.label("svc_cat"))
         .select_from(c.outerjoin(s, s.c.id == c.c.service_id)).where(c.c.encounter_id == aid, c.c.deleted_at.is_(None)))
    if scope == "procedures":
        q = q.where(sa.not_(sa.func.coalesce(c.c.source, "").like("ipd.daily%")), sa.func.coalesce(c.c.source, "").notin_(list(PHARM_SOURCES)),
                    sa.not_(sa.and_(c.c.source == ORD, sa.func.coalesce(s.c.category, "").in_(["LAB", "RAD"]))),
                    sa.func.coalesce(s.c.category, "") != "PHARM")
    rows = (await ctx.conn.execute(q.order_by(c.c.charged_at.desc().nulls_last()).limit(limit))).all()
    users_ = await _user_names(ctx, [r.created_by for r in rows])
    money_ok = _money_ok(ctx)
    out = []
    for r in rows:
        d = r.data or {}
        typ = SOURCE_TYPE.get(r.source or "", "Room / daily" if (r.source or "").startswith("ipd.daily") else "Charge")
        if r.svc_cat == "PACKAGE":
            typ = "Package"
        elif r.source == ORD and r.svc_cat == "PROC":
            typ = "Procedure"
        x = {"id": str(r.id), "code": r.code, "at": _iso(r.charged_at), "description": d.get("description") or r.svc_name, "service": r.svc_name,
             "category": r.svc_cat, "type": typ, "source": r.source, "qty": float(r.qty or 0),
             "status": {"Unbilled": "Charged"}.get(r.status or "", r.status), "grade": d.get("grade"),
             "performed_by": d.get("performed_by_name"), "entered_by": d.get("entered_by") or users_.get(r.created_by or ""),
             "voided_by": d.get("voided_by"), "voided_at": d.get("voided_at"), "void_reason": d.get("void_reason"),
             "can_void": r.status == "Unbilled" and _perm(ctx, "void_charge") and r.source not in (ORD, USAGE)}
        if money_ok:
            x.update(rate=float(r.rate or 0), discount=float(r.discount or 0), amount=float(r.amount or 0))
        out.append(x)
    return out


@router.get("/admissions/{aid}/charges")
async def charge_list(aid: str, ctx: CtxDep, scope: Literal["procedures", "all"] = "procedures"):
    rs.ensure_access(ctx, E(ADM), "read")
    await _admission(ctx, aid)
    rows = await _charge_rows(ctx, aid, scope=scope)
    out: dict[str, Any] = {"rows": rows}
    if _money_ok(ctx):
        live_rows = [r for r in rows if r["status"] != "Cancelled"]
        out["total"] = round(sum(r.get("amount") or 0 for r in live_rows), 2)
    return out


# ================================================================== cancellations / void

class VoidIn(BaseModel):
    target: Literal["order", "charge", "usage", "indent"]
    id: str
    reason: str = Field(min_length=3, max_length=500)


@router.post("/admissions/{aid}/void")
async def void(aid: str, body: VoidIn, ctx: CtxDep):
    await _admission(ctx, aid)
    reason = body.reason.strip()
    stamp = {"voided_by": _who(ctx), "voided_at": _stamp(), "void_reason": reason}
    if body.target == "order":
        _need(ctx, "void_order", "cancel orders")
        o = await get(ctx, ORD, body.id)
        if not o or o.get("encounter_id") != aid:
            raise not_found("Order not found on this admission")
        if o.get("status") not in ("Ordered", "Collected"):
            raise conflict(f"{o.get('code')} is {o.get('status')} — it can no longer be cancelled here")
        await transition(ctx, ORD, body.id, "Cancelled", note=reason, values={"cancel_reason": reason})
        rec = await update(ctx, ORD, body.id, stamp)
        return {"id": rec["id"], "status": rec.get("status"), **stamp}
    if body.target == "charge":
        _need(ctx, "void_charge", "cancel charges")
        ch = await get(ctx, CHARGES, body.id)
        if not ch or ch.get("encounter_id") != aid:
            raise not_found("Charge not found on this admission")
        if ch.get("status") != "Unbilled":
            raise conflict("Only charges not yet billed can be cancelled — raise a credit note for billed charges")
        if ch.get("source") in (ORD, USAGE):
            raise conflict("Cancel the order / equipment usage instead — its charge is reversed with it")
        rec = await update(ctx, CHARGES, body.id, {"status": "Cancelled", **stamp})
        await audit(ctx, "charge.void", module="billing", entity="charges", record=rec, note=reason)
        return {"id": rec["id"], "status": rec.get("status"), **stamp}
    if body.target == "usage":
        _need(ctx, "void_usage", "cancel equipment usage")
        return await _void_usage(ctx, aid, body.id, stamp)
    _need(ctx, "void_indent", "cancel indents")
    ind = await get(ctx, INDENT, body.id)
    if not ind or ind.get("encounter_id") != aid:
        raise not_found("Indent not found on this admission")
    if ind.get("status") != "Raised":
        raise conflict(f"{ind.get('code')} is already {str(ind.get('status')).lower()}")
    await transition(ctx, INDENT, body.id, "Cancelled", note=reason)
    rec = await update(ctx, INDENT, body.id, stamp)
    return {"id": rec["id"], "status": rec.get("status"), **stamp}


# ================================================================== medication: search, advised, indents, returns, ledger

async def _stock_by_item(ctx, item_ids: list[str], store: str) -> dict[str, float]:
    from ..schema import stock

    ids = [i for i in item_ids if is_uuid(i)]
    if not ids:
        return {}
    t = stock
    rows = (await ctx.conn.execute(sa.select(t.c.item_id, sa.func.coalesce(sa.func.sum(t.c.qty), 0)).where(
        t.c.item_id.in_(ids), t.c.store == store, t.c.deleted_at.is_(None), t.c.qty > 0,
        sa.func.coalesce(t.c.status, "Available") != "Quarantine", sa.or_(t.c.expiry.is_(None), t.c.expiry >= now().date()))
        .group_by(t.c.item_id))).all()
    return {str(i): float(q) for i, q in rows}


def _packing(it: dict) -> str:
    return " · ".join(str(x) for x in (it.get("form"), it.get("strength"), it.get("uom")) if x) or "—"


def _med(it: dict, avail: float | None, money_ok: bool) -> dict:
    x = {"id": it["id"], "code": it.get("code"), "name": item_label(it), "generic": it.get("generic_name"), "form": it.get("form"),
         "strength": it.get("strength"), "uom": it.get("uom"), "packing": _packing(it), "available": avail, "schedule": it.get("schedule"),
         "high_alert": bool(it.get("high_alert")), "type": it.get("type")}
    if money_ok:
        x["mrp"] = it.get("mrp")
    return x


@router.get("/medicines")
async def medicine_search(ctx: CtxDep, q: str = "", generic: bool = False, store: str = "IP Pharmacy", limit: int = 20):
    rs.ensure_access(ctx, E(INDENT), "read")
    if store not in INDENT_STORES:
        raise bad_request("Unknown stock point")
    term = q.strip().lower()
    if len(term) < 2:
        return {"rows": []}
    ent = E("mdm.items")
    t = ent.table
    like = f"%{term}%"
    conds = [*rs.base_where(ent), t.c.data["type"].astext.in_(["Drug", "Consumable", "Implant"]),
             sa.func.coalesce(t.c.status, "Active") != "Inactive"]
    match = [sa.func.lower(t.c.data["generic_name"].astext).like(like)] if generic else [
        sa.func.lower(t.c.title).like(like), sa.func.lower(t.c.code).like(like), sa.func.lower(t.c.data["generic_name"].astext).like(like)]
    rows = (await ctx.conn.execute(sa.select(t).where(*conds, sa.or_(*match)).order_by(t.c.title).limit(max(1, min(limit, 50))))).all()
    items = [rs.to_record(r) for r in rows]
    stock_ = await _stock_by_item(ctx, [i["id"] for i in items], store)
    money_ok = _money_ok(ctx)
    return {"rows": [_med(i, stock_.get(i["id"], 0.0), money_ok) for i in items], "store": store}


@router.get("/admissions/{aid}/advised-medications")
async def advised_medications(aid: str, ctx: CtxDep, store: str = "IP Pharmacy"):
    """The doctor's active prescriptions for this patient, with IP-pharmacy availability, to pull into an indent."""
    rs.ensure_access(ctx, E(INDENT), "read")
    adm = await _admission(ctx, aid)
    meds = await active_medications(ctx, adm["patient_id"])
    items = await items_by_id(ctx, [m["item_id"] for m in meds])
    stock_ = await _stock_by_item(ctx, list(items), store if store in INDENT_STORES else "IP Pharmacy")
    money_ok = _money_ok(ctx)
    out = []
    for m in meds:
        it = items.get(m["item_id"])
        if not it:
            continue
        out.append({**_med(it, stock_.get(it["id"], 0.0), money_ok), "dose": m.get("dose"), "route": m.get("route"),
                    "frequency": m.get("frequency"), "duration_days": m.get("duration_days"), "suggested_qty": m.get("qty") or 1,
                    "prescription": m.get("prescription_code"), "prescribed_at": m.get("prescribed_at")})
    return {"rows": out}


class IndentLine(BaseModel):
    item_id: str
    qty: float = Field(gt=0, le=10000)


class IndentIn(BaseModel):
    category: Literal["Normal", "Urgent", "STAT", "Discharge medication"] = "Normal"
    store: str = "IP Pharmacy"
    items: list[IndentLine] = Field(min_length=1, max_length=40)
    remarks: str | None = Field(default=None, max_length=500)
    override_reason: str | None = Field(default=None, max_length=500)


@router.post("/admissions/{aid}/indents", status_code=201)
async def raise_indent(aid: str, body: IndentIn, ctx: CtxDep):
    """Patient-specific ward indent to the IP pharmacy, screened for allergy / interaction / duplicate therapy
    first. The pharmacy issues it (stock out + charge to this stay) from its own worklist."""
    _need(ctx, "indent", "raise medication indents")
    adm = await _admission(ctx, aid)
    _require_open(adm, "raise indents")
    if body.store not in INDENT_STORES:
        raise bad_request("Unknown stock point")
    ids = [ln.item_id for ln in body.items]
    if len(set(ids)) != len(ids):
        raise bad_request("A medicine is listed twice — combine the quantities")
    items = await items_by_id(ctx, ids)
    bad = [i for i in ids if i not in items or items[i].get("type") not in ("Drug", "Consumable", "Implant")]
    if bad:
        raise bad_request("Not an indentable medicine / consumable")
    bed = await _bed_info(ctx, adm.get("bed_id"))
    if not bed or not bed.get("ward_id"):
        raise conflict("The patient has no ward bed — indents go to a ward")
    lines = [{"item_id": ln.item_id, "qty": ln.qty} for ln in body.items]
    safety = await safety_check(ctx, patient_id=adm["patient_id"], lines=lines, purpose="indent", store=body.store)
    alerts = safety["alerts"]
    blocking = [a for a in alerts if a["blocking"]]
    if blocking:
        raise ApiError(409, blocking[0]["message"], {"alerts": alerts}, code="INDENT_BLOCKED")
    severe = [a for a in alerts if a["severity"] in ("Major", "Contraindicated") and a["type"] != "stock"]
    reason = (body.override_reason or "").strip()
    if severe and len(reason) < 5:
        raise ApiError(409, f"{len(severe)} serious alert(s) need an override reason", {"alerts": alerts}, code="OVERRIDE_REQUIRED")
    controlled = any(items[i].get("schedule") in ("X", "Narcotic") for i in ids)
    rec = await create(ctx, INDENT, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "ward": bed["ward_id"], "requested_by": _emp(ctx),
        "indent_type": "Narcotic / Schedule X" if controlled else "Patient-specific", "priority": INDENT_CATEGORIES[body.category],
        "required_by": _stamp(), "items": [{"item": ln.item_id, "qty": ln.qty} for ln in body.items], "remarks": body.remarks,
        "indent_category": body.category, "store": body.store, "raised_by": _who(ctx), "bed_no": bed.get("bed_no"),
        "alerts_overridden": "\n".join([f"[{a['severity']}] {a['message']}" for a in severe] + [f"Override reason: {reason}"]) if severe else None,
        "status": "Raised",
    })
    if severe:
        await audit(ctx, "pharmacy.override", module="pharmacy", entity="indents", record=rec,
                    after={"alerts": [a["message"] for a in severe]}, note=reason)
    live(ctx, "alerts", {"type": "ward_indent", "roles": ["pharmacist"], "indent_id": rec["id"],
                         "message": f"{body.category} indent {rec.get('code')} · {bed.get('ward')} bed {bed.get('bed_no')} · {len(lines)} item(s)"})
    return {"indent": {"id": rec["id"], "code": rec.get("code"), "status": rec.get("status")}, "alerts": alerts}


@router.get("/admissions/{aid}/indents")
async def indents(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(INDENT), "read")
    await _admission(ctx, aid)
    ti = E(INDENT).table
    rows = await _rows(ctx, INDENT, ti.c.encounter_id == aid, order=ti.c.created_at.desc(), limit=200)
    items = await items_by_id(ctx, [ln.get("item") for r in rows for ln in r.get("items") or []])
    return {"rows": [{"id": r["id"], "code": r.get("code"), "status": r.get("status"), "category": r.get("indent_category") or r.get("priority"),
                      "priority": r.get("priority"), "store": r.get("store") or "IP Pharmacy", "raised_by": r.get("raised_by"),
                      "raised_at": r.get("created_at"), "issued_at": r.get("issued_at"), "remarks": r.get("remarks"),
                      "voided_by": r.get("voided_by"), "void_reason": r.get("void_reason"),
                      "can_cancel": r.get("status") == "Raised" and _perm(ctx, "void_indent"),
                      "items": [{"item_id": ln.get("item"), "name": item_label(items.get(ln.get("item"))), "qty": ln.get("qty"),
                                 "issued_qty": ln.get("issued_qty"), "packing": _packing(items.get(ln.get("item")) or {})}
                                for ln in r.get("items") or []]} for r in rows]}


@router.get("/admissions/{aid}/returnable")
async def returnable(aid: str, ctx: CtxDep):
    """Medicines dispensed on this stay that can still be returned (per dispense, item and batch)."""
    rs.ensure_access(ctx, E(RETURN), "read")
    await _admission(ctx, aid)
    dt = E(DISPENSE).table
    disps = await _rows(ctx, DISPENSE, dt.c.encounter_id == aid, dt.c.status.in_(["Dispensed", "Partially Returned"]),
                        order=dt.c.created_at.desc(), limit=100)
    rt = E(RETURN).table
    rets = await _rows(ctx, RETURN, rt.c.encounter_id == aid, rt.c.status.in_(["Requested", "Accepted"]), limit=300)
    back: dict[tuple[str, str, str | None], float] = defaultdict(float)
    for r in rets:
        for ln in r.get("items") or []:
            back[(r.get("dispense"), ln.get("item"), ln.get("batch"))] += float(ln.get("qty") or 0)
    items = await items_by_id(ctx, [ln.get("item") for d in disps for ln in d.get("items") or []])
    money_ok = _money_ok(ctx)
    out = []
    for d in disps:
        agg: dict[tuple[str, str | None], dict] = {}
        for ln in d.get("items") or []:
            k = (ln.get("item"), ln.get("batch"))
            a = agg.setdefault(k, {"qty": 0.0, "rate": ln.get("rate")})
            a["qty"] += float(ln.get("qty") or 0)
        for (iid, batch), a in agg.items():
            left = a["qty"] - back.get((d["id"], iid, batch), 0.0)
            if left <= 0:
                continue
            x = {"dispense_id": d["id"], "dispense": d.get("code"), "store": d.get("store"), "dispensed_at": d.get("dispensed_at"),
                 "item_id": iid, "name": item_label(items.get(iid)), "batch": batch, "dispensed_qty": a["qty"], "returnable_qty": left}
            if money_ok:
                x["rate"] = a["rate"]
            out.append(x)
    returns = [{"id": r["id"], "code": r.get("code"), "status": r.get("status"), "reason": r.get("reason"), "at": r.get("created_at"),
                "dispense": r.get("dispense"), "items": len(r.get("items") or []),
                **({"refund_amount": r.get("refund_amount")} if money_ok else {})} for r in sorted(rets, key=lambda r: str(r.get("created_at")), reverse=True)]
    return {"rows": out, "returns": returns, "reasons": E(RETURN).field_map["reason"].options}


class ReturnLine(BaseModel):
    item_id: str
    batch: str | None = None
    qty: float = Field(gt=0, le=10000)


class ReturnIn(BaseModel):
    dispense_id: str
    items: list[ReturnLine] = Field(min_length=1, max_length=40)
    reason: str


@router.post("/admissions/{aid}/returns", status_code=201)
async def pharmacy_return(aid: str, body: ReturnIn, ctx: CtxDep):
    _need(ctx, "pharmacy_return", "return medicines to the pharmacy")
    adm = await _admission(ctx, aid)
    disp = await get(ctx, DISPENSE, body.dispense_id)
    if not disp or disp.get("encounter_id") != aid:
        raise not_found("Dispense not found on this admission")
    if body.reason not in (E(RETURN).field_map["reason"].options or []):
        raise bad_request("Choose the return reason", {"fields": {"reason": "Required"}})
    bed = await _bed_info(ctx, adm.get("bed_id"))
    rec = await create(ctx, RETURN, {"patient_id": adm["patient_id"], "encounter_id": aid, "dispense": body.dispense_id,
                                     "items": [{"item": ln.item_id, "batch": ln.batch, "qty": ln.qty} for ln in body.items],
                                     "reason": body.reason, "ward": (bed or {}).get("ward_id"), "returned_by": _emp(ctx), "status": "Requested"})
    live(ctx, "alerts", {"type": "pharmacy_return", "roles": ["pharmacist"],
                         "message": f"Return {rec.get('code')} from {(bed or {}).get('ward') or 'ward'} awaiting acceptance"})
    out = {"id": rec["id"], "code": rec.get("code"), "status": rec.get("status")}
    if _money_ok(ctx):
        out["refund_amount"] = rec.get("refund_amount")
    return out


_DESC = re.compile(r"^(?:Return:\s*)?(?P<name>.*?)\s+·\s+batch\s+(?P<batch>\S+)", re.I)


@router.get("/admissions/{aid}/medicines")
async def medicines_ledger(aid: str, ctx: CtxDep, q: str | None = None):
    """Everything medicinal billed to this stay — ward indents, IP-pharmacy direct billing, the pharmacy counter
    and returns (negative lines) — with batch, quantity, amount, source and reference."""
    if not (_ent_ok(ctx, INDENT) or _ent_ok(ctx, DISPENSE) or _money_ok(ctx)):
        raise forbidden("Your role cannot view the medicines ledger")
    await _admission(ctx, aid)
    c, s = charges, services
    rows = (await ctx.conn.execute(
        sa.select(c.c.id, c.c.status, c.c.qty, c.c.rate, c.c.amount, c.c.source, c.c.source_id, c.c.charged_at, c.c.data,
                  s.c.category.label("cat"))
        .select_from(c.outerjoin(s, s.c.id == c.c.service_id))
        .where(c.c.encounter_id == aid, c.c.deleted_at.is_(None), c.c.status != "Cancelled",
               sa.or_(c.c.source.in_(list(PHARM_SOURCES)), s.c.category == "PHARM"))
        .order_by(c.c.charged_at.desc()))).all()
    src_ids = list({str(r.source_id) for r in rows if r.source_id})
    refs: dict[str, dict] = {}
    if src_ids:
        for r in (await ctx.conn.execute(sa.select(records.c.id, records.c.code, records.c.data).where(records.c.id.in_(src_ids)))).all():
            refs[str(r.id)] = {"code": r.code, "store": (r.data or {}).get("store")}
    moves: dict[tuple[str, str | None], dict] = {}
    if src_ids:
        m = stock_moves
        for mv in (await ctx.conn.execute(sa.select(m.c.item_id, m.c.batch, m.c.store, m.c.ref, m.c.data).where(
                m.c.data["ref_id"].astext.in_(src_ids)))).all():
            moves.setdefault(((mv.data or {}).get("ref_id"), mv.batch), {"item_id": str(mv.item_id), "store": mv.store})
    its = await items_by_id(ctx, [v["item_id"] for v in moves.values()])
    money_ok = _money_ok(ctx)
    out = []
    for r in rows:
        d = r.data or {}
        desc = d.get("description") or ""
        mm = _DESC.match(desc)
        name, batch = (mm.group("name"), mm.group("batch")) if mm else (desc, None)
        if batch == "—":
            batch = None
        sid = str(r.source_id) if r.source_id else None
        mv = moves.get((sid, batch)) or {}
        it = its.get(mv.get("item_id") or "")
        ref = refs.get(sid or "") or {}
        if r.source == INDENT:
            label, kind = "Ward indent", "indent"
        elif r.source == RETURN:
            label, kind = "Returned", "return"
        elif r.source == DISPENSE:
            store = ref.get("store") or mv.get("store")
            label, kind = ("Pharmacy counter", "counter") if store == "OP Pharmacy" else ("Direct billing", "direct")
        else:
            label, kind = "Added on bill", "bill"
        x = {"id": str(r.id), "at": _iso(r.charged_at), "medicine": name.replace(" (ward indent)", ""), "item_code": (it or {}).get("code"),
             "batch": batch, "qty": float(r.qty or 0), "source": label, "kind": kind, "reference": ref.get("code"),
             "store": ref.get("store") or mv.get("store") or ("IP Pharmacy" if r.source == INDENT else None), "billed": r.status == "Billed"}
        if money_ok:
            x["amount"] = float(r.amount or 0)
        out.append(x)
    if q and q.strip():
        term = q.strip().lower()
        out = [x for x in out if term in (x["medicine"] or "").lower() or term in (x["item_code"] or "").lower() or term in (x["batch"] or "").lower()]
    res: dict[str, Any] = {"rows": out}
    if money_ok:
        res["totals"] = {"issued": round(sum(x["amount"] for x in out if x["amount"] > 0), 2),
                         "returned": round(-sum(x["amount"] for x in out if x["amount"] < 0), 2),
                         "net": round(sum(x["amount"] for x in out), 2)}
    return res


# ================================================================== equipment usage billing

@router.get("/equipment")
async def equipment_master(ctx: CtxDep, q: str | None = None, status: str | None = None):
    rs.ensure_access(ctx, E(EQUIP), "read")
    rows = await find(ctx, EQUIP, q=q, status=status, limit=300, sort="title")
    money_ok = _money_ok(ctx)
    return {"rows": [{"id": r["id"], "code": r.get("code"), "name": r.get("name"), "tag": r.get("tag"), "type": r.get("equipment_type"),
                      "billing_rule": r.get("billing_rule"), "status": r.get("status"), "current_admission": r.get("current_admission"),
                      **({"rate": r.get("rate")} if money_ok else {})} for r in rows]}


def usage_units(rule: str, start: datetime, end: datetime, min_units: float | None) -> tuple[int, float]:
    hours = max(0.0, (end - start).total_seconds() / 3600)
    floor_ = int(min_units) if min_units not in (None, "") and float(min_units) > 0 else 1
    if rule == "Per use":
        return 1, round(hours, 2)
    units = math.ceil(round(hours / 24 if rule == "Per day" else hours, 6))
    return max(floor_, units), round(hours, 2)


@router.get("/admissions/{aid}/equipment")
async def admission_equipment(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(USAGE), "read")
    await _admission(ctx, aid)
    ut = E(USAGE).table
    rows = await _rows(ctx, USAGE, ut.c.encounter_id == aid, order=_dcol(USAGE, "connected_at").desc(), limit=200)
    money_ok = _money_ok(ctx)
    ts = now()
    out = []
    for u in rows:
        x = {"id": u["id"], "code": u.get("code"), "equipment_id": u.get("equipment"), "name": u.get("equipment_name"), "tag": u.get("equipment_tag"),
             "type": u.get("equipment_type"), "billing_rule": u.get("billing_rule"), "connected_at": u.get("connected_at"),
             "disconnected_at": u.get("disconnected_at"), "connected_by": u.get("connected_by"), "disconnected_by": u.get("disconnected_by"),
             "units": u.get("units"), "hours": u.get("hours"), "status": u.get("status"), "remarks": u.get("remarks"),
             "voided_by": u.get("voided_by"), "void_reason": u.get("void_reason")}
        if u.get("status") == "Connected":
            units, hrs = usage_units(u.get("billing_rule") or "Per hour", parse_dt(u["connected_at"]), ts, u.get("min_units"))
            x.update(running_units=units, running_hours=hrs)
            if money_ok:
                x["running_amount"] = float(money(Decimal(str(u.get("rate") or 0)) * units))
        if money_ok:
            x.update(rate=u.get("rate"), amount=u.get("amount"))
        out.append(x)
    avail = await find(ctx, EQUIP, status="Available", limit=300, sort="title")
    return {"active": [x for x in out if x["status"] == "Connected"], "history": [x for x in out if x["status"] != "Connected"],
            "available": [{"id": e["id"], "name": e.get("name"), "tag": e.get("tag"), "type": e.get("equipment_type"),
                           "billing_rule": e.get("billing_rule"), **({"rate": e.get("rate")} if money_ok else {})} for e in avail],
            "configured": await _count(ctx, EQUIP)}


class ConnectIn(BaseModel):
    equipment_id: str | None = None
    tag: str | None = Field(default=None, max_length=80)
    connected_at: str | None = None
    remarks: str | None = Field(default=None, max_length=300)


@router.post("/admissions/{aid}/equipment/connect", status_code=201)
async def connect_equipment(aid: str, body: ConnectIn, ctx: CtxDep):
    _need(ctx, "equipment", "connect equipment")
    adm = await _admission(ctx, aid)
    _require_open(adm, "connect equipment")
    eq = None
    if body.equipment_id:
        eq = await get(ctx, EQUIP, body.equipment_id)
    elif body.tag and body.tag.strip():
        t = E(EQUIP).table
        tag = body.tag.strip().lower()
        row = (await ctx.conn.execute(sa.select(t.c.id).where(*rs.base_where(E(EQUIP)), sa.or_(
            sa.func.lower(t.c.data["tag"].astext) == tag, sa.func.lower(t.c.code) == tag)).limit(1))).first()
        eq = await get(ctx, EQUIP, str(row.id)) if row else None
    if not eq:
        raise not_found("No equipment with that tag — ask billing to add it to the equipment master")
    if eq.get("status") != "Available":
        where = ""
        if eq.get("current_admission"):
            other = await get(ctx, ADM, eq["current_admission"])
            where = f" on {other.get('code')}" if other else ""
        raise conflict(f"{eq.get('name')} is {str(eq.get('status')).lower()}{where}")
    at = parse_dt(body.connected_at) if body.connected_at else now()
    if at > now() + timedelta(minutes=5):
        raise bad_request("'Connected at' cannot be in the future", {"fields": {"connected_at": "Future time"}})
    started = parse_dt(adm.get("started_at"))
    if started and at < started:
        raise bad_request("'Connected at' is before the admission", {"fields": {"connected_at": "Before admission"}})
    rec = await create(ctx, USAGE, {
        "patient_id": adm["patient_id"], "encounter_id": aid, "equipment": eq["id"], "equipment_name": eq.get("name"),
        "equipment_tag": eq.get("tag"), "equipment_type": eq.get("equipment_type"), "connected_at": at.isoformat(timespec="minutes"),
        "billing_rule": eq.get("billing_rule"), "rate": eq.get("rate"), "min_units": eq.get("min_units"), "service": eq.get("service"),
        "connected_by": _who(ctx), "remarks": (body.remarks or "").strip() or None, "status": "Connected",
    })
    await update(ctx, EQUIP, eq["id"], {"status": "In Use", "current_admission": aid, "connected_at": at.isoformat(timespec="minutes")})
    return {"id": rec["id"], "code": rec.get("code"), "name": eq.get("name"), "connected_at": rec.get("connected_at")}


class DisconnectIn(BaseModel):
    disconnected_at: str | None = None
    remarks: str | None = Field(default=None, max_length=300)


@router.post("/equipment-usage/{uid}/disconnect")
async def disconnect_equipment(uid: str, body: DisconnectIn, ctx: CtxDep):
    """Stop the machine's clock and post its charge by the rate rule (per hour / per day rounded up, or per use)."""
    _need(ctx, "equipment", "disconnect equipment")
    u = await get(ctx, USAGE, uid)
    if not u:
        raise not_found("Equipment usage not found")
    if u.get("status") != "Connected":
        raise conflict(f"Already {str(u.get('status')).lower()}")
    start = parse_dt(u["connected_at"])
    end = parse_dt(body.disconnected_at) if body.disconnected_at else now()
    if end > now() + timedelta(minutes=5):
        raise bad_request("'Disconnected at' cannot be in the future", {"fields": {"disconnected_at": "Future time"}})
    if end < start:
        raise bad_request("'Disconnected at' is before it was connected", {"fields": {"disconnected_at": "Before connection"}})
    rule = u.get("billing_rule") or "Per hour"
    units, hours = usage_units(rule, start, end, u.get("min_units"))
    rate = Decimal(str(u.get("rate") or 0))
    unit_word = {"Per hour": "h", "Per day": "day(s)", "Per use": "use"}.get(rule, "unit(s)")
    ch = await post_charge(ctx, patient_id=u["patient_id"], encounter_id=u.get("encounter_id"), service_id=u.get("service"), qty=units,
                           rate=float(rate), description=f"{u.get('equipment_name')} — {units} {unit_word} "
                                                         f"({start.astimezone(settings.tz):%d %b %H:%M} → {end.astimezone(settings.tz):%d %b %H:%M})",
                           source=USAGE, source_id=uid)
    rec = await update(ctx, USAGE, uid, {"disconnected_at": end.isoformat(timespec="minutes"), "units": units, "hours": hours,
                                         "amount": float(money(rate * units)), "charge": ch["id"], "disconnected_by": _who(ctx),
                                         "status": "Disconnected", "remarks": " · ".join(x for x in (u.get("remarks"), (body.remarks or "").strip()) if x) or None})
    await update(ctx, EQUIP, u["equipment"], {"status": "Available", "current_admission": None, "connected_at": None})
    out = {"id": rec["id"], "units": units, "hours": hours, "billing_rule": rule, "charge": {"id": ch["id"], "code": ch.get("code")}}
    if _money_ok(ctx):
        out["amount"] = rec.get("amount")
    return out


async def _void_usage(ctx, aid: str, uid: str, stamp: dict) -> dict:
    u = await get(ctx, USAGE, uid)
    if not u or u.get("encounter_id") != aid:
        raise not_found("Equipment usage not found on this admission")
    if u.get("status") == "Voided":
        raise conflict("Already voided")
    if u.get("status") == "Disconnected" and u.get("charge"):
        ch = await get(ctx, CHARGES, u["charge"])
        if ch and ch.get("status") == "Billed":
            raise conflict("The equipment charge is already billed — raise a credit note in billing")
        if ch and ch.get("status") == "Unbilled":
            await update(ctx, CHARGES, ch["id"], {"status": "Cancelled", **stamp})
            await audit(ctx, "charge.void", module="billing", entity="charges", record=ch, note=stamp["void_reason"])
    rec = await update(ctx, USAGE, uid, {"status": "Voided", **stamp})
    if u.get("status") == "Connected":
        await update(ctx, EQUIP, u["equipment"], {"status": "Available", "current_admission": None, "connected_at": None})
    return {"id": rec["id"], "status": rec.get("status"), **stamp}


# ================================================================== On Bed census

@router.get("/on-bed")
async def on_bed(ctx: CtxDep, ward: str | None = None, floor: str | None = None, q: str | None = None):
    """Every occupied bed right now, with stay, consultant, acuity and (for billing roles) the running money position:
    billed = all live charges of the stay; collected = payments on the stay's bills + the patient's unused advance."""
    rs.ensure_access(ctx, E(ADM), "read")
    e, p, b, doc = encounters, patients, beds, employees
    w = records.alias("w")
    dp = records.alias("dp")
    stmt = (sa.select(e.c.id, e.c.code, e.c.status, e.c.patient_id, e.c.started_at, e.c.data, b.c.id.label("bed_id"), b.c.bed_no,
                      b.c.category, b.c.ward_id, w.c.title.label("ward"), w.c.data.label("wdata"), p.c.code.label("uhid"), p.c.first_name,
                      p.c.last_name, p.c.gender, p.c.dob, doc.c.name.label("doctor"), dp.c.title.label("doctor_dept"))
            .select_from(e.join(b, b.c.id == e.c.bed_id).join(p, p.c.id == e.c.patient_id).outerjoin(w, w.c.id == b.c.ward_id)
                         .outerjoin(doc, doc.c.id == e.c.doctor_id).outerjoin(dp, dp.c.id == doc.c.department_id))
            .where(e.c.type == "IPD", e.c.status.in_(OPEN), e.c.deleted_at.is_(None), b.c.status == "Occupied"))
    if ward:
        if not is_uuid(ward):
            raise bad_request("Invalid ward")
        stmt = stmt.where(b.c.ward_id == ward)
    if floor:
        stmt = stmt.where(w.c.data["floor"].astext == floor)
    if q and q.strip():
        term = f"%{q.strip().lower()}%"
        stmt = stmt.where(sa.or_(sa.func.lower(sa.func.concat(p.c.first_name, " ", p.c.last_name)).like(term), sa.func.lower(p.c.code).like(term),
                                 sa.func.lower(e.c.code).like(term), sa.func.lower(b.c.bed_no).like(term), sa.func.lower(doc.c.name).like(term),
                                 sa.func.lower(w.c.title).like(term)))
    rows = (await ctx.conn.execute(stmt.order_by(w.c.title, b.c.bed_no))).all()
    seen: set[str] = set()
    rows = [r for r in rows if not (str(r.bed_id) in seen or seen.add(str(r.bed_id)))]  # legacy overlaps: one row per bed
    money_ok = _money_ok(ctx)
    billed: dict[str, Decimal] = {}
    paid: dict[str, Decimal] = {}
    dep: dict[str, Decimal] = {}
    if money_ok and rows:
        enc_ids = [str(r.id) for r in rows]
        pids = list({str(r.patient_id) for r in rows})
        c = charges
        billed = {str(k): Decimal(v) for k, v in (await ctx.conn.execute(sa.select(c.c.encounter_id, sa.func.coalesce(sa.func.sum(c.c.amount), 0))
                                                                         .where(c.c.encounter_id.in_(enc_ids), c.c.deleted_at.is_(None), c.c.status != "Cancelled")
                                                                         .group_by(c.c.encounter_id))).all()}
        pay, inv = payments, invoices
        signed = sa.case((pay.c.kind == "REFUND", -pay.c.amount), else_=pay.c.amount)
        paid = {str(k): Decimal(v) for k, v in (await ctx.conn.execute(
            sa.select(inv.c.encounter_id, sa.func.coalesce(sa.func.sum(signed), 0)).select_from(pay.join(inv, inv.c.id == pay.c.invoice_id))
            .where(inv.c.encounter_id.in_(enc_ids), pay.c.deleted_at.is_(None), pay.c.status == "Success",
                   pay.c.kind.in_(["PAYMENT", "ADJUSTMENT", "REFUND"]), pay.c.mode != "DEPOSIT")
            .group_by(inv.c.encounter_id))).all()}
        dsigned = sa.case((pay.c.kind == "DEPOSIT", pay.c.amount), else_=-pay.c.amount)
        dep = {str(k): Decimal(v) for k, v in (await ctx.conn.execute(
            sa.select(pay.c.patient_id, sa.func.coalesce(sa.func.sum(dsigned), 0)).where(
                pay.c.patient_id.in_(pids), pay.c.deleted_at.is_(None), pay.c.status == "Success", sa.or_(
                    pay.c.kind == "DEPOSIT", sa.and_(pay.c.kind == "ADJUSTMENT", pay.c.mode == "DEPOSIT"),
                    sa.and_(pay.c.kind == "REFUND", pay.c.invoice_id.is_(None)))).group_by(pay.c.patient_id))).all()}
        # deposits used on this stay's bills were counted as ADJUSTMENT payments above
        used_here = {str(k): Decimal(v) for k, v in (await ctx.conn.execute(
            sa.select(inv.c.encounter_id, sa.func.coalesce(sa.func.sum(pay.c.amount), 0)).select_from(pay.join(inv, inv.c.id == pay.c.invoice_id))
            .where(inv.c.encounter_id.in_(enc_ids), pay.c.deleted_at.is_(None), pay.c.status == "Success", pay.c.kind == "ADJUSTMENT",
                   pay.c.mode == "DEPOSIT").group_by(inv.c.encounter_id))).all()}
        for k, v in used_here.items():
            paid[k] = paid.get(k, Decimal(0)) + v
    ts = now()
    out = []
    tot = {"billed": Decimal(0), "collected": Decimal(0), "due": Decimal(0), "advance": Decimal(0)}
    floors: set[str] = set()
    for r in rows:
        d = r.data or {}
        wd = r.wdata or {}
        if wd.get("floor") not in (None, ""):
            floors.add(str(wd.get("floor")))
        x = {"id": str(r.id), "code": r.code, "status": r.status, "discharging": r.status == "Discharge Initiated",
             "bed": {"id": str(r.bed_id), "bed_no": r.bed_no, "category": r.category},
             "ward": {"id": str(r.ward_id) if r.ward_id else None, "name": r.ward, "floor": wd.get("floor"), "block": wd.get("block")},
             "patient": {"id": str(r.patient_id), "name": _name(r.first_name, r.last_name), "uhid": r.uhid, "gender": r.gender, "age": age_from(r.dob)},
             "started_at": _iso(r.started_at), "los_days": _los_days(r.started_at), "doctor": r.doctor, "doctor_department": r.doctor_dept,
             "condition": d.get("condition") or "Not assessed", "payer": d.get("payer") or "Cash", "isolation": d.get("isolation"),
             "accepted": bool(d.get("accepted_at"))}
        if money_ok:
            bl = billed.get(x["id"], Decimal(0))
            col = paid.get(x["id"], Decimal(0)) + max(dep.get(x["patient"]["id"], Decimal(0)), Decimal(0))
            bal = bl - col
            x.update(billed=float(money(bl)), collected=float(money(col)), balance=float(money(bal)))
            tot["billed"] += bl
            tot["collected"] += col
            tot["due" if bal > 0 else "advance"] += abs(bal)
        out.append(x)
    wards = await _ward_rows(ctx)
    res: dict[str, Any] = {
        "generated_at": ts.isoformat(timespec="seconds"), "rows": out, "count": len(out), "money": money_ok,
        "kpis": {"on_bed": len(out), "discharging": sum(1 for x in out if x["discharging"]),
                 "critical": sum(1 for x in out if x["condition"] == "Critical"), "awaiting_acceptance": sum(1 for x in out if not x["accepted"])},
        "filters": {"wards": [{"id": w_["id"], "name": w_["name"], "floor": w_["floor"]} for w_ in wards],
                    "floors": sorted({str(w_["floor"]) for w_ in wards if w_.get("floor") not in (None, "")} | floors,
                                     key=lambda f: (0, int(f)) if f.isdigit() else (1, f))},
    }
    res["can"] = {"tabs": _perms(ctx)["tabs"], "bill": money_ok and _ent_ok(ctx, "billing.invoices")}
    if money_ok:
        res["kpis"].update(billed=float(money(tot["billed"])), collected=float(money(tot["collected"])),
                           balance_to_collect=float(money(tot["due"])), advance_held=float(money(tot["advance"])))
    return res


# ================================================================== discharge tracker, clearances, step notifications

async def _log_notification(ctx, aid: str, entry: dict) -> None:
    """Append to the admission's notification log without a record-service write (no audit noise)."""
    cur = (await ctx.conn.execute(sa.select(encounters.c.data["discharge_notifications"]).where(encounters.c.id == aid))).scalar() or []
    lst = [*(cur if isinstance(cur, list) else []), entry][-60:]
    await ctx.conn.execute(encounters.update().where(encounters.c.id == aid).values(
        data=sa.func.coalesce(encounters.c.data, sa.text("'{}'::jsonb")).op("||")(
            sa.func.jsonb_build_object("discharge_notifications", sa.cast(sa.literal(lst, type_=encounters.c.data.type), encounters.c.data.type)))))


async def notify_step(ctx, adm: dict, step: str, message: str) -> list[str]:
    """Tell the roles configured for this discharge-tracker step (per-admission override → settings → defaults):
    a realtime alert to those roles, the patient-facing notifications engine (only if an active template exists for
    `ipd.<step>`), and an entry in the admission's notification log."""
    cfg = await discharge_settings(ctx)
    roles = list((adm.get("discharge_notify") or {}).get(step, cfg["notify"].get(step, [])))
    if roles:
        live(ctx, "alerts", {"type": "discharge_step", "step": step, "roles": roles, "admission_id": adm["id"],
                             "patient_id": adm.get("patient_id"), "message": message})
    try:
        from .notify import notify as patient_notify

        await patient_notify(ctx, f"ipd.{step}", values={"admission_no": adm.get("code"), "step": dict(STEPS).get(step, step)},
                             patient_id=adm.get("patient_id"), ref_id=adm["id"])
    except Exception:  # never block a clinical action on a messaging problem
        log.exception("patient notification for ipd.%s failed", step)
    await _log_notification(ctx, adm["id"], {"step": step, "roles": roles, "at": _stamp(), "by": _who(ctx), "message": message})
    return roles


async def ensure_clearances(ctx, adm: dict) -> list[dict]:
    cfg = await discharge_settings(ctx)
    ct = E(CLR).table
    have = {c.get("department_type"): c for c in await _rows(ctx, CLR, ct.c.encounter_id == adm["id"], limit=50)}
    bill = await clearance(ctx, adm["id"])
    out = []
    for dept in cfg["clearances"]:
        if dept in have:
            out.append(have[dept])
            continue
        due = float(bill.get("balance_due") or 0) + float(bill.get("unbilled_amount") or 0)
        out.append(await create(ctx, CLR, {"patient_id": adm["patient_id"], "encounter_id": adm["id"], "department_type": dept,
                                           "pending_amount": due if dept in ("Billing", "Insurance / TPA") and due > 0 else None,
                                           "status": "Pending"}))
    return out


async def _clearance_blocker(ctx, adm: dict, dept: str, bill: dict | None = None) -> str | None:
    if dept in ("Billing", "Insurance / TPA"):
        bill = bill or await clearance(ctx, adm["id"])
        if not bill["cleared"]:
            return "Final bill not settled"
    if dept == "Pharmacy returns" and has(RETURN):
        rt = E(RETURN).table
        n = await _count(ctx, RETURN, rt.c.patient_id == adm["patient_id"], rt.c.status == "Requested")
        if n:
            return f"{n} pharmacy return(s) awaiting acceptance"
    return None


async def tracker_state(ctx, adm: dict) -> dict:
    aid = adm["id"]
    cfg = await discharge_settings(ctx)
    notify = {k: list((adm.get("discharge_notify") or {}).get(k, cfg["notify"].get(k, []))) for k in STEP_KEYS}
    ct = E(CLR).table
    clr = await _rows(ctx, CLR, ct.c.encounter_id == aid, order=ct.c.created_at.asc(), limit=50)
    bill = await clearance(ctx, aid)
    inv_n = (await ctx.conn.execute(sa.select(sa.func.count()).where(
        invoices.c.encounter_id == aid, invoices.c.deleted_at.is_(None), invoices.c.status != "Cancelled"))).scalar_one()
    status = adm.get("status")
    discharged = status == "Discharged"
    initiated = bool(adm.get("discharge_initiated_at")) or discharged
    cleared = [c for c in clr if c.get("status") == "Cleared"]
    clr_done = bool(clr) and len(cleared) == len(clr)
    settled = bool(bill["cleared"]) and (inv_n > 0 or discharged)
    money_ok = _money_ok(ctx)
    due = float(bill.get("balance_due") or 0) + float(bill.get("unbilled_amount") or 0)
    open_ = status in OPEN
    depts = []
    for c in clr:
        dept = c.get("department_type")
        roles = sorted(CLEARANCE_ROLES.get(dept, {"nurse"}))
        blocker = await _clearance_blocker(ctx, adm, dept, bill) if c.get("status") != "Cleared" else None
        mine = _can(ctx, *roles)
        depts.append({"id": c["id"], "code": c.get("code"), "department": dept, "status": c.get("status"), "cleared_at": c.get("cleared_at"),
                      "cleared_by": c.get("cleared_by"), "hold_reason": c.get("hold_reason"), "remarks": c.get("remarks"), "roles": roles,
                      "blocker": blocker, "can_sign": mine and open_ and c.get("status") != "Cleared",
                      "can_hold": mine and open_ and c.get("status") == "Pending"})
    steps = [
        {"key": "admitted", "label": "Admitted", "done": True, "at": adm.get("started_at"), "by": adm.get("admitted_by"),
         "detail": adm.get("code")},
        {"key": "initiated", "label": "Discharge initiated", "done": initiated, "at": adm.get("discharge_initiated_at"),
         "by": adm.get("discharge_initiated_by"),
         "detail": (adm.get("discharge_type") or "Routine") if initiated else "Awaiting initiation from the chart or the discharge desk"},
        {"key": "clearances", "label": "Clearances", "done": clr_done,
         "at": max((c.get("cleared_at") for c in cleared if c.get("cleared_at")), default=None) if clr_done else None,
         "detail": f"{len(cleared)} of {len(clr)} departments cleared" if clr else "Departments are asked once discharge is initiated"},
        {"key": "settlement", "label": "Settlement", "done": settled, "at": None,
         "detail": "Settled" if settled else ((f"₹{due:,.2f} to settle" if money_ok else "Bill not settled") if due > 0 or inv_n
                                              else "Final bill not generated yet")},
        {"key": "discharged", "label": "Discharged", "done": discharged, "at": adm.get("ended_at") if discharged else None,
         "detail": adm.get("discharge_type") if discharged else "Pending"},
    ]
    nxt = next((s for s in steps if not s["done"]), None)
    for s in steps:
        partial = (s["key"] == "clearances" and bool(cleared) and not clr_done) or (s["key"] == "settlement" and inv_n > 0 and not settled)
        s["state"] = "done" if s["done"] else "in_progress" if s is nxt or partial else "pending"
        s["notify_roles"] = notify[s["key"]]
    res = {"admission": {"id": aid, "code": adm.get("code"), "status": status}, "steps": steps, "done": sum(1 for s in steps if s["done"]),
           "total": len(steps), "next": nxt["label"] if nxt else None, "clearances": depts, "configured_clearances": cfg["clearances"],
           "notifications": list(reversed((adm.get("discharge_notifications") or [])[-20:])), "roles": STAFF_ROLES,
           "billing": {"cleared": bill["cleared"], "invoices": inv_n,
                       **({"balance_due": bill.get("balance_due"), "unbilled_amount": bill.get("unbilled_amount")} if money_ok else {})},
           "can_configure": _perm(ctx, "configure") and open_, "can_initiate": _perm(ctx, "discharge") and status == "Admitted",
           "can_revoke": _perm(ctx, "revoke_discharge") and status == "Discharge Initiated"}
    return res


@router.get("/admissions/{aid}/tracker")
async def tracker(aid: str, ctx: CtxDep):
    rs.ensure_access(ctx, E(ADM), "read")
    adm = await _admission(ctx, aid)
    if adm.get("status") == "Discharge Initiated" and not await _count(ctx, CLR, E(CLR).table.c.encounter_id == aid):
        await ensure_clearances(ctx, adm)  # admissions initiated before departmental clearances existed
    return await tracker_state(ctx, adm)


class NotifyRolesIn(BaseModel):
    step: Literal["admitted", "initiated", "clearances", "settlement", "discharged"]
    roles: list[str] = Field(default_factory=list, max_length=12)


@router.put("/admissions/{aid}/tracker/notify")
async def set_notify_roles(aid: str, body: NotifyRolesIn, ctx: CtxDep):
    _need(ctx, "configure", "change who is notified")
    adm = await _admission(ctx, aid)
    _require_open(adm, "change notifications")
    bad = [r for r in body.roles if r not in STAFF_ROLES]
    if bad:
        raise bad_request(f"Unknown role: {', '.join(bad)}")
    cur = dict(adm.get("discharge_notify") or {})
    cur[body.step] = sorted(set(body.roles))
    await update(ctx, ADM, aid, {"discharge_notify": cur})
    return await tracker_state(ctx, await _admission(ctx, aid))


class ClearIn(BaseModel):
    remarks: str | None = Field(default=None, max_length=300)


class HoldIn(BaseModel):
    reason: str = Field(min_length=3, max_length=300)


async def _clearance_for(ctx, cid: str) -> tuple[dict, dict]:
    c = await get(ctx, CLR, cid)
    if not c:
        raise not_found("Clearance not found")
    adm = await _admission(ctx, c.get("encounter_id"))
    return c, adm


@router.post("/clearances/{cid}/clear")
async def sign_clearance(cid: str, body: ClearIn, ctx: CtxDep):
    c, adm = await _clearance_for(ctx, cid)
    _require_open(adm, "sign clearances")
    if c.get("status") == "Cleared":
        raise conflict(f"{c.get('department_type')} is already cleared by {c.get('cleared_by')}")
    await transition(ctx, CLR, cid, "Cleared", values={"remarks": body.remarks} if (body.remarks or "").strip() else None)
    return await tracker_state(ctx, await _admission(ctx, adm["id"]))


@router.post("/clearances/{cid}/hold")
async def hold_clearance(cid: str, body: HoldIn, ctx: CtxDep):
    c, adm = await _clearance_for(ctx, cid)
    _require_open(adm, "update clearances")
    if not _can(ctx, *CLEARANCE_ROLES.get(c.get("department_type"), {"nurse"})):
        raise forbidden(f"{c.get('department_type')} clearance is signed by {', '.join(sorted(CLEARANCE_ROLES.get(c.get('department_type'), {'nurse'})))}")
    if c.get("status") != "Pending":
        raise conflict(f"Clearance is {str(c.get('status')).lower()}")
    await transition(ctx, CLR, cid, "On Hold", values={"hold_reason": body.reason.strip()})
    return await tracker_state(ctx, await _admission(ctx, adm["id"]))


@router.get("/clearances/pending")
async def pending_clearances(ctx: CtxDep):
    """Department queue: discharge clearances waiting for the caller's role (pharmacy, lab, radiology, billing …)."""
    depts = [d for d, roles in CLEARANCE_ROLES.items() if _can(ctx, *roles)]
    if not depts:
        raise forbidden("Your role signs no discharge clearances")
    ent = E(CLR)
    t = ent.table
    e, p, b = encounters, patients, beds
    w = records.alias("w")
    rows = (await ctx.conn.execute(
        sa.select(t.c.id, t.c.code, t.c.status, t.c.created_at, t.c.data, e.c.id.label("aid"), e.c.code.label("adm_code"),
                  e.c.status.label("adm_status"), e.c.data.label("adata"), p.c.code.label("uhid"), p.c.first_name, p.c.last_name,
                  b.c.bed_no, w.c.title.label("ward"))
        .select_from(t.join(e, e.c.id == t.c.encounter_id).join(p, p.c.id == e.c.patient_id).outerjoin(b, b.c.id == e.c.bed_id)
                     .outerjoin(w, w.c.id == b.c.ward_id))
        .where(*rs.base_where(ent), t.c.status.in_(["Pending", "On Hold"]), t.c.data["department_type"].astext.in_(depts),
               e.c.status == "Discharge Initiated", e.c.deleted_at.is_(None))
        .order_by(t.c.created_at))).all()
    ts = now()
    out = []
    for r in rows:
        d, ad = r.data or {}, r.adata or {}
        out.append({"id": str(r.id), "code": r.code, "department": d.get("department_type"), "status": r.status,
                    "hold_reason": d.get("hold_reason"), "waiting_minutes": minutes_between(r.created_at, ts),
                    "admission": {"id": str(r.aid), "code": r.adm_code, "discharge_type": ad.get("discharge_type"),
                                  "initiated_at": ad.get("discharge_initiated_at")},
                    "patient": {"name": _name(r.first_name, r.last_name), "uhid": r.uhid}, "bed": r.bed_no, "ward": r.ward})
    return {"rows": out, "departments": depts}


@hook(CLR, "before_transition")
async def _clearance_rules(ctx, nxt: dict, prev: dict) -> None:
    if nxt["status"] != "Cleared":
        return
    dept = nxt.get("department_type")
    roles = CLEARANCE_ROLES.get(dept, {"nurse"})
    if ctx.user_id != "system" and not _can(ctx, *roles):
        raise forbidden(f"{dept} clearance is signed off by {', '.join(sorted(roles))}")
    adm = await get(ctx, ADM, nxt.get("encounter_id")) if nxt.get("encounter_id") else None
    if adm:
        why = await _clearance_blocker(ctx, adm, dept)
        if why:
            raise conflict(f"Cannot clear {dept}: {why}")
    nxt["cleared_by"] = f"{_who(ctx)} ({ctx.role})"


@hook(CLR, "after_transition")
async def _clearance_done(ctx, rec: dict, prev: dict) -> None:
    if rec.get("status") != "Cleared" or not rec.get("encounter_id"):
        return
    item = CLEARANCE_CHECKLIST.get(rec.get("department_type"))
    if item:
        chk = await first(ctx, CHK, encounter_id=rec["encounter_id"], status="Open")
        if chk and not chk.get(item):
            try:
                await update(ctx, CHK, chk["id"], {item: True})
            except ApiError as e:  # the checklist enforces its own sign-off roles
                log.info("checklist %s not ticked from clearance: %s", item, e.detail)
    ct = E(CLR).table
    rows = await _rows(ctx, CLR, ct.c.encounter_id == rec["encounter_id"], limit=50)
    if rows and all(r.get("status") == "Cleared" for r in rows):
        adm = await get(ctx, ADM, rec["encounter_id"])
        if adm:
            await notify_step(ctx, adm, "clearances", f"All departmental clearances done for {adm.get('code')} — ready for settlement")


@hook(ADM, "before_transition")
async def _equipment_before_discharge(ctx, nxt: dict, prev: dict) -> None:
    if nxt["status"] not in ("Discharged", "Cancelled"):
        return
    ut = E(USAGE).table
    n = await _count(ctx, USAGE, ut.c.encounter_id == prev["id"], ut.c.status == "Connected")
    if n:
        raise conflict(f"{n} machine(s) are still connected to this patient — disconnect them first (usage is billed on disconnect)")


@hook(ADM, "after_transition")
async def _discharge_steps(ctx, rec: dict, prev: dict) -> None:
    to = rec.get("status")
    bed = await _bed_info(ctx, rec.get("bed_id"))
    where = f"{bed['ward']} bed {bed['bed_no']}" if bed else "ward"
    if to == "Discharge Initiated":
        await ensure_clearances(ctx, rec)
        await notify_step(ctx, rec, "initiated", f"Discharge initiated for {rec.get('code')} ({where}) — {rec.get('discharge_type') or 'Routine'}. "
                                                  "Departments: please complete your clearance.")
    elif to == "Discharged":
        await notify_step(ctx, rec, "discharged", f"{rec.get('code')} discharged — {where} released for cleaning")


@hook("billing.invoices", "after_update")
async def _settlement_watch(ctx, rec: dict, prev: dict) -> None:
    eid = rec.get("encounter_id")
    if not eid or float(rec.get("balance") or 0) > 0 or float(prev.get("balance") or 0) <= 0:
        return
    adm = await get(ctx, ADM, eid)
    if not adm or adm.get("status") != "Discharge Initiated":
        return
    if any(n.get("step") == "settlement" for n in adm.get("discharge_notifications") or []):
        return
    if (await clearance(ctx, eid))["cleared"]:
        await notify_step(ctx, adm, "settlement", f"Bill settled for {adm.get('code')} — patient can be discharged")


# ================================================================== print documents (wristband, face sheet, case summary)

@router.get("/admissions/{aid}/print")
async def print_data(aid: str, ctx: CtxDep, doc: Literal["wristband", "face-sheet", "case-summary"] = "wristband"):
    rs.ensure_access(ctx, E(ADM), "read")
    if doc == "case-summary" and not _can(ctx, *CLINICAL_VIEW):
        raise forbidden("Your role cannot print the case summary")
    hdr = await chart_header(aid, ctx)
    adm = await _admission(ctx, aid)
    out: dict[str, Any] = {"doc": doc, "hospital": {"name": settings.hospital_name}, "printed_at": _stamp(), "printed_by": _who(ctx),
                           **{k: hdr[k] for k in ("admission", "patient", "bed", "department", "doctor", "second_doctor", "allergies", "reactive")}}
    if doc == "case-summary":
        out["vitals"] = hdr["last_obs"]
        o = orders
        rows = (await ctx.conn.execute(_order_q().where(o.c.encounter_id == aid, o.c.status != "Cancelled").order_by(o.c.ordered_at))).all()
        out["investigations"] = [{"test": r.svc_name or r.svc_code, "status": _stage(r.status), "ordered_at": _iso(r.ordered_at),
                                  "reported_at": _iso(r.verified_at or r.resulted_at), "critical": bool(r.critical),
                                  "abnormal": [f"{v.get('parameter')} {v.get('value')}{' ' + v['unit'] if v.get('unit') else ''} ({v.get('flag')})"
                                               for v in (r.data or {}).get("result_values") or [] if v.get("flag") not in (None, "Normal")][:6],
                                  "impression": (r.data or {}).get("impression")} for r in rows]
        meds = await active_medications(ctx, adm["patient_id"])
        its = await items_by_id(ctx, [m["item_id"] for m in meds])
        out["medications"] = [{"name": item_label(its.get(m["item_id"])), "dose": m.get("dose"), "route": m.get("route"),
                               "frequency": m.get("frequency")} for m in meds]
        out["procedures"] = [{k: x.get(k) for k in ("at", "description", "type", "status")} for x in await _charge_rows(ctx, aid, scope="procedures")
                             if x["status"] != "Cancelled" and x["type"] in ("Procedure", "Package", "Equipment")]
        nt = E(NOTES).table
        nrows = await _rows(ctx, NOTES, nt.c.encounter_id == aid, nt.c.status != "Draft", order=_dcol(NOTES, "noted_at").asc(), limit=60)
        an = await _emp_names(ctx, [n.get("author") for n in nrows])
        out["notes"] = [{"noted_at": n.get("noted_at"), "type": n.get("note_type"), "shift": n.get("shift"), "author": an.get(n.get("author") or ""),
                         "note": n.get("note")} for n in nrows if n.get("note_type") != "Nursing Note"][-15:]
        summ = await first(ctx, SUMMARY, encounter_id=aid, sort="-created_at")
        out["discharge_summary"] = {k: summ.get(k) for k in ("code", "status", "final_diagnosis", "condition_at_discharge", "follow_up")} if summ else None
    await audit(ctx, "print", module="ipd", entity="admissions", record=adm, note=doc)
    return out


# ================================================================== list exports (Excel / CSV / JSON / print)

def _export_admissions(rows: list[dict]) -> tuple[list[tuple[str, str]], list[dict]]:
    cols = [("code", "Admission no"), ("patient", "Patient"), ("uhid", "UHID"), ("age_sex", "Age / sex"), ("phone", "Mobile"),
            ("ward", "Ward"), ("bed", "Bed"), ("category", "Category"), ("sponsor", "Sponsor"), ("admitted", "Admitted"),
            ("consultant", "Consultant"), ("diagnosis", "Provisional diagnosis"), ("condition", "Condition"), ("status", "Status"),
            ("accepted_by", "Ward acceptance"), ("los", "Stay (days)")]
    out = [{"code": r["code"], "patient": r["patient"]["name"], "uhid": r["patient"]["uhid"],
            "age_sex": " ".join(str(x) for x in (r["patient"].get("age"), (r["patient"].get("gender") or "")[:1]) if x not in (None, "")),
            "phone": r["patient"].get("phone"), "ward": r["ward"]["name"], "bed": r["bed_no"], "category": r["payer_label"],
            "sponsor": (r["sponsor"] or {}).get("name") or "Self-pay", "admitted": r["started_at"], "consultant": r["doctor"],
            "diagnosis": r["diagnosis"], "condition": r["condition"], "status": r["status"],
            "accepted_by": f"{r['accepted_by']} · {r['accepted_at'][:16].replace('T', ' ')}" if r.get("accepted_at") else "", "los": r["los_days"]}
           for r in rows]
    return cols, out


def _export_on_bed(rows: list[dict], money_ok: bool) -> tuple[list[tuple[str, str]], list[dict]]:
    cols = [("ward", "Ward"), ("floor", "Floor"), ("bed", "Bed"), ("patient", "Patient"), ("uhid", "UHID"), ("code", "IP no"),
            ("status", "Status"), ("stay", "Stay (days)"), ("since", "Admitted"), ("consultant", "Consultant"), ("department", "Department"),
            ("acuity", "Acuity")]
    if money_ok:
        cols += [("billed", "Bill total"), ("collected", "Collected"), ("balance", "Balance")]
    out = [{"ward": r["ward"]["name"], "floor": r["ward"].get("floor"), "bed": r["bed"]["bed_no"], "patient": r["patient"]["name"],
            "uhid": r["patient"]["uhid"], "code": r["code"], "status": "Discharging" if r["discharging"] else r["status"], "stay": r["los_days"],
            "since": r["started_at"], "consultant": r["doctor"], "department": r.get("doctor_department"), "acuity": r["condition"],
            **({"billed": r.get("billed"), "collected": r.get("collected"), "balance": r.get("balance")} if money_ok else {})} for r in rows]
    return cols, out


@router.get("/lists/{name}/export.{ext}")
async def export_list(name: Literal["admissions", "on-bed"], ext: Literal["csv", "json", "xlsx"], ctx: CtxDep, request: Request):
    qp = request.query_params
    f = {k: qp.get(k) or None for k in ("q", "status", "ward", "sponsor", "acceptance", "condition", "date_from", "date_to", "floor")}
    if name == "admissions":
        data = await admissions_data(ctx, **{k: v for k, v in f.items() if k != "floor"}, sort=qp.get("sort") or "-started_at",
                                     page=1, limit=10_000, cap=10_000)
        cols, rows = _export_admissions(data["rows"])
        label, total = "IP admissions", data["total"]
    else:
        data = await on_bed(ctx, ward=f["ward"], floor=f["floor"], q=f["q"])
        cols, rows = _export_on_bed(data["rows"], data["money"])
        label, total = "On bed census", data["count"]
    wanted = [c for c in (qp.get("cols") or "").split(",") if c]
    if wanted:
        cols = [c for c in cols if c[0] in wanted] or cols
    await audit(ctx, "export", module="ipd", entity=name, note=f"{ext} · {len(rows)} rows")
    stamp = now().strftime("%Y%m%d")
    disposition = "inline" if qp.get("inline") else "attachment"
    if ext == "json":
        return JSONResponse({"label": label, "total": total, "exported_at": _stamp(), "columns": [{"key": k, "label": lb} for k, lb in cols],
                             "rows": [{k: r.get(k) for k, _ in cols} for r in rows]},
                            headers={"Content-Disposition": f'{disposition}; filename="ipd-{name}-{stamp}.json"'})
    if ext == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([lb for _, lb in cols])
        for r in rows:
            w.writerow(["" if r.get(k) is None else r.get(k) for k, _ in cols])
        return Response("\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="ipd-{name}-{stamp}.csv"'})
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = label[:31]
    ws.append([lb for _, lb in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F4C5C")
    for r in rows:
        ws.append([r.get(k) for k, _ in cols])
    for i, (k, lb) in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = min(max(10, max([len(lb)] + [len(str(r.get(k) or "")) for r in rows[:200]]) + 2), 50)
        if k in ("billed", "collected", "balance"):
            for cell in ws[letter][1:]:
                cell.number_format = "#,##,##0.00"
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="ipd-{name}-{stamp}.xlsx"'})


__all__ = ["router", "vital_status", "usage_units", "visit_tracker", "tracker_state", "notify_step", "discharge_settings", "PERMS"]
