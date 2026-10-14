"""Generic REST API for every module entity: /api/m/{module}/{entity}."""

from __future__ import annotations

import io
from typing import Any

from fastapi import APIRouter, Request
from fastapi.responses import JSONResponse, PlainTextResponse, Response
from pydantic import BaseModel

from ..core import records as svc
from ..core.audit import audit, audit_trail
from ..core.config_store import effective_fields
from ..core.rbac import check_access
from ..core.registry import Entity, get_module
from ..core.security import CtxDep
from ..core.util import ApiError, forbidden, now, to_csv

router = APIRouter(prefix="/api/m", tags=["records"])


def _filters(request: Request) -> dict[str, Any]:
    """`f_<field>=value` query params → equality filters (comma = any of)."""
    out = {}
    for k, v in request.query_params.items():
        if k.startswith("f_") and v != "":
            out[k[2:]] = v.split(",") if "," in v else v
    return out


@router.get("/{module}/stats")
async def module_stats(module: str, ctx: CtxDep):
    mod = get_module(module)
    if not mod:
        raise ApiError(404, f"Unknown module {module}")
    if not check_access(ctx.user, mod, None, "read").ok and not any(
        check_access(ctx.user, mod, e, "read").ok for e in mod.entities
    ):
        raise forbidden()
    ents = {}
    for e in mod.entities:
        if check_access(ctx.user, mod, e, "read").ok and not e.restricted:
            ents[e.key] = await svc.entity_stats(ctx, e)
    return {"entities": ents, "kpis": await svc.compute_kpis(ctx, mod)}


@router.get("/{module}/{entity}")
async def list_(module: str, entity: str, request: Request, ctx: CtxDep, q: str | None = None,
                status: str | None = None, patient_id: str | None = None, encounter_id: str | None = None,
                employee_id: str | None = None, date_from: str | None = None, date_to: str | None = None,
                sort: str | None = None, page: int = 1, limit: int = 25):
    ent = svc.entity_or_404(module, entity)
    return await svc.list_records(ctx, ent, q=q, status=status, patient_id=patient_id, encounter_id=encounter_id,
                                  employee_id=employee_id, filters=_filters(request), date_from=date_from,
                                  date_to=date_to, sort=sort, page=page, limit=limit)


EXPORT_MAX_ROWS = 10_000
BASE_COLS = {"code": "Code", "status": "Status", "title": "Title", "patient_id": "Patient", "employee_id": "Employee",
             "encounter_id": "Encounter", "created_at": "Created", "updated_at": "Updated"}


def mask_value(v: Any) -> Any:
    """Show only the last 4 characters (mobile numbers, ID numbers) — DPDP-friendly exports and lists."""
    if v in (None, ""):
        return v
    s = str(v)
    return "•" * len(s) if len(s) <= 4 else "X" * (len(s) - 4) + s[-4:]


async def _export_rows(ctx, ent: Entity, request: Request, *, q, status, date_from, date_to, sort, cols: str | None,
                       patient_id: str | None = None, encounter_id: str | None = None, employee_id: str | None = None):
    """Rows (refs resolved to labels, configured masks applied) and (key, label) columns for exports."""
    fc, _ = await effective_fields(ctx, ent.fqn)
    rows: list[dict] = []
    page, total = 1, 0
    while len(rows) < EXPORT_MAX_ROWS:
        res = await svc.list_records(ctx, ent, q=q, status=status, patient_id=patient_id, encounter_id=encounter_id,
                                     employee_id=employee_id, filters=_filters(request), date_from=date_from,
                                     date_to=date_to, sort=sort, page=page, limit=2000, max_limit=2000)
        total = res["total"]
        refs = res["refs"]
        for r in res["rows"]:
            rows.append({k: refs.get(v, v) if isinstance(v, str) else v for k, v in r.items()})
        if not res["rows"] or len(rows) >= total:
            break
        page += 1
    labels = {f.key: (fc.get(f.key) or {}).get("label") or f.label for f in ent.fields}
    hidden = {k for k, c in fc.items() if c.get("visible") is False}
    if cols:
        keys = [c for c in cols.split(",") if c in BASE_COLS or c in ent.field_map]
        columns = [(k, BASE_COLS.get(k) or labels.get(k, k)) for k in keys]
    else:
        columns = [("code", "Code")] + [(f.key, labels[f.key]) for f in ent.fields
                                        if f.type not in ("lines", "signature", "file") and not f.hidden and f.key not in hidden] + [
            ("status", "Status"), ("created_at", "Created")]
    masked = {k for k, c in fc.items() if c.get("mask")}
    for r in rows:
        for k in masked:
            if k in r:
                r[k] = mask_value(r[k])
        for k, v in list(r.items()):
            if isinstance(v, list):
                r[k] = ", ".join(str(x) for x in v if not isinstance(x, dict)) if not any(isinstance(x, dict) for x in v) else f"{len(v)} lines"
            elif isinstance(v, bool):
                r[k] = "Yes" if v else "No"
    return rows[:EXPORT_MAX_ROWS], columns, total


def _export_args(request: Request) -> dict:
    qp = request.query_params
    return {k: qp.get(k) for k in ("q", "status", "date_from", "date_to", "sort", "cols", "patient_id", "encounter_id", "employee_id")}


@router.get("/{module}/{entity}/export.csv", response_class=PlainTextResponse)
async def export_csv(module: str, entity: str, request: Request, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rows, cols, total = await _export_rows(ctx, ent, request, **_export_args(request))
    await audit(ctx, "export", module=ent.module_id, entity=ent.key, note=f"CSV · {len(rows)} of {total} rows")
    return PlainTextResponse("\ufeff" + to_csv(rows, cols), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f'attachment; filename="{ent.fqn}-{now():%Y%m%d}.csv"'})


@router.get("/{module}/{entity}/export.json")
async def export_json(module: str, entity: str, request: Request, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rows, cols, total = await _export_rows(ctx, ent, request, **_export_args(request))
    await audit(ctx, "export", module=ent.module_id, entity=ent.key, note=f"JSON · {len(rows)} of {total} rows")
    body = {"entity": ent.fqn, "label": ent.label, "exported_at": now().isoformat(timespec="seconds"), "total": total,
            "columns": [{"key": k, "label": label} for k, label in cols],
            "rows": [{k: r.get(k) for k, _ in cols} for r in rows]}
    disposition = "inline" if request.query_params.get("inline") else "attachment"
    return JSONResponse(body, headers={"Content-Disposition": f'{disposition}; filename="{ent.fqn}-{now():%Y%m%d}.json"'})


@router.get("/{module}/{entity}/export.xlsx")
async def export_xlsx(module: str, entity: str, request: Request, ctx: CtxDep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ent = svc.entity_or_404(module, entity)
    rows, cols, total = await _export_rows(ctx, ent, request, **_export_args(request))
    wb = Workbook()
    ws = wb.active
    ws.title = ent.label[:31]
    ws.append([label for _, label in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F4C5C")
        c.alignment = Alignment(vertical="center")
    money = {f.key for f in ent.fields if f.type == "currency"}
    for r in rows:
        ws.append([_cell(r.get(k)) for k, _ in cols])
    for i, (k, label) in enumerate(cols, start=1):
        letter = get_column_letter(i)
        width = max([len(str(label))] + [len(str(r.get(k) or "")) for r in rows[:200]])
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
        if k in money:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    await audit(ctx, "export", module=ent.module_id, entity=ent.key, note=f"Excel · {len(rows)} of {total} rows")
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{ent.fqn}-{now():%Y%m%d}.xlsx"'})


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v)
    if len(s) >= 10 and s[4:5] == "-" and s[7:8] == "-":  # ISO dates → readable, Excel-sortable text
        return s[:16].replace("T", " ")
    return s[:32000]


@router.get("/{module}/{entity}/{rid}")
async def read(module: str, entity: str, rid: str, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rec = await svc.get_record(ctx, ent, rid)
    rec["_transitions"] = svc.allowed_transitions(ent, rec, ctx.role)
    return rec


@router.post("/{module}/{entity}", status_code=201)
async def create(module: str, entity: str, body: dict[str, Any], ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rec = await svc.create_record(ctx, ent, body)
    rec["_refs"] = await svc.resolve_refs(ctx, ent, [rec])
    return rec


@router.put("/{module}/{entity}/{rid}")
async def update(module: str, entity: str, rid: str, body: dict[str, Any], ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rec = await svc.update_record(ctx, ent, rid, body)
    rec["_refs"] = await svc.resolve_refs(ctx, ent, [rec])
    return rec


class TransitionIn(BaseModel):
    to: str
    note: str | None = None
    values: dict[str, Any] | None = None


@router.post("/{module}/{entity}/{rid}/transition")
async def transition(module: str, entity: str, rid: str, body: TransitionIn, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rec = await svc.transition_record(ctx, ent, rid, body.to, note=body.note, values=body.values)
    rec["_refs"] = await svc.resolve_refs(ctx, ent, [rec])
    rec["_transitions"] = svc.allowed_transitions(ent, rec, ctx.role)
    return rec


@router.delete("/{module}/{entity}/{rid}")
async def delete(module: str, entity: str, rid: str, ctx: CtxDep, reason: str | None = None):
    ent = svc.entity_or_404(module, entity)
    if ctx.role not in ("admin",) and not (ent.roles or ent.module.roles):
        raise forbidden()
    await svc.delete_record(ctx, ent, rid, reason=reason)
    return {"ok": True}


@router.get("/{module}/{entity}/{rid}/audit")
async def record_audit(module: str, entity: str, rid: str, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    svc.ensure_access(ctx, ent, "read")
    return await audit_trail(ctx.conn, record_id=rid, limit=200)


@router.post("/{module}/{entity}/{rid}/print")
async def record_print(module: str, entity: str, rid: str, ctx: CtxDep):
    ent = svc.entity_or_404(module, entity)
    rec = await svc.get_record(ctx, ent, rid)
    await audit(ctx, "print", module=ent.module_id, entity=ent.key, record=rec)
    return {"ok": True}
